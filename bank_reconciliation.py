import os
import sys
import re
import io
import json
import argparse
import warnings
from datetime import datetime, timedelta

CONFIG_PATH = os.path.join(os.path.dirname(__file__), "config.json")
CREDENTIALS_PATH = os.path.join(os.path.dirname(__file__), "credentials.json")
OUTPUT_DIR = os.path.join(os.path.dirname(__file__), "output")

def load_config():
    with open(CONFIG_PATH, "r", encoding="utf-8") as f:
        return json.load(f)

def get_drive_headers():
    from google.oauth2.service_account import Credentials
    from google.auth.transport.requests import Request
    scopes = ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive"]
    creds = Credentials.from_service_account_file(CREDENTIALS_PATH, scopes=scopes)
    creds.refresh(Request())
    return {"Authorization": f"Bearer {creds.token}"}, creds

def to_roc_date(date_str):
    clean = str(date_str).replace("/", "-").strip()
    parts = clean.split("-")
    if len(parts) == 3 and parts[0].isdigit():
        y = int(parts[0]) - 1911
        m = int(parts[1])
        d = int(parts[2])
        return f"{y:03d}/{m:02d}/{d:02d}", f"{y:03d}{m:02d}{d:02d}"
    return clean, clean.replace("/", "")

OFFICIAL_HEADER_1 = [
    "傳票號碼", "傳票種類", "傳票日期", "傳票總摘要", "序號", "借貸", "會計項目代號",
    "部門代號", "專案代號", "對象類別", "對象編號", "幣別", "匯率", "原幣金額", "本幣金額", "摘要",
    "建立人員", "建立日期", "主檔自定欄位1", "主檔自定欄位2", "主檔自定欄位3", "主檔自定欄位4",
    "主檔自定欄位5", "主檔自定欄位6", "主檔自定欄位7", "主檔自定欄位8", "主檔自定欄位9",
    "主檔自定欄位10", "主檔自定欄位11", "主檔自定欄位12", "明細自定欄位1", "明細自定欄位2",
    "明細自定欄位3", "明細自定欄位4", "明細自定欄位5", "明細自定欄位6", "明細自定欄位7",
    "明細自定欄位8", "明細自定欄位9", "明細自定欄位10", "明細自定欄位11", "明細自定欄位12"
]
OFFICIAL_HEADER_2 = [
    "DTGL004", "DTGL006", "DTGL007", "MFGL006", "DTGL005", "DTGL012", "DTGL008",
    "DTGL014", "DTGL015", "DTGL100", "DTGL009", "DTGL027", "DTGL028", "DTGL021", "DTGL013", "DTGL011",
    "MFGL996", "MFGL997", "MFGL801", "MFGL802", "MFGL803", "MFGL804",
    "MFGL805", "MFGL806", "MFGL807", "MFGL808", "MFGL809", "MFGL810", "MFGL811", "MFGL812",
    "DTGL801", "DTGL802", "DTGL803", "DTGL804", "DTGL805", "DTGL806", "DTGL807", "DTGL808",
    "DTGL809", "DTGL810", "DTGL811", "DTGL812"
]

def format_official_42_row(v):
    amt = round(v.get("本幣金額", 0))
    dc = 1 if str(v.get("借貸", "借")) in ("借", "1") else 2
    d_std = v.get("傳票日期_標準")
    if not d_std:
        raw_d = str(v.get("傳票日期", "")).split()[0].replace("/", "-")
        parts = raw_d.split("-")
        if len(parts) == 3 and parts[0].isdigit():
            y = int(parts[0]) + 1911 if int(parts[0]) < 1900 else int(parts[0])
            m = int(parts[1])
            d = int(parts[2])
            d_std = f"{y:04d}-{m:02d}-{d:02d}"
        else:
            d_std = f"{raw_d}"
    else:
        d_std = str(d_std).split()[0]
    row = [
        v.get("傳票號碼", "@"),
        3, # 轉帳傳票
        d_std,
        v.get("總摘要", ""),
        int(v.get("項次", 1)),
        dc,
        v.get("會計項目", ""),
        v.get("部門", ""),
        v.get("專案", ""),
        "01",
        "A01",
        "NTD",
        1,
        amt,
        amt,
        v.get("摘要", "")
    ]
    row.extend([""] * 26)
    return row

def reconcile_store_bank(store_name, store_cfg, target_month="2026-10", preview=False):
    import requests, openpyxl, gspread
    headers, creds = get_drive_headers()
    gc = gspread.authorize(creds)

    bank_file_id = store_cfg.get("esun_bank_file_id")
    if not bank_file_id:
        print(f"\nℹ️ 【{store_name}】尚未設定玉山銀行試算表 ID，暫時略過。")
        return []

    print("\n" + "="*70)
    print(f"🏦 啟動【{store_name}】玉山銀行智慧對帳與沖銷傳票產生")
    print(f"📄 對帳月份: 【{target_month}】")
    print("="*70)

    # 取得設定帳號代碼
    accounts = store_cfg.get("accounts", {})
    bank_acct_code = "1102"
    bank_acct_name = "銀行存款-玉山"
    fee_acct_code = "5151"
    fee_acct_name = "手續費支出"
    fee_tax_code = "1268"
    fee_tax_name = "進項稅額"

    cash_code = accounts.get("cash", {}).get("code", "1130-01")
    cash_name = accounts.get("cash", {}).get("name", "應收帳款-現金")

    linepay_code = accounts.get("linepay", {}).get("code", "1130-02")
    linepay_name = accounts.get("linepay", {}).get("name", "應收帳款-linepay")

    ubereats_code = accounts.get("ubereats", {}).get("code", "1130-03")
    ubereats_name = accounts.get("ubereats", {}).get("name", "應收帳款-ubereats")

    foodpanda_code = accounts.get("foodpanda", {}).get("code", "1130-04")
    foodpanda_name = accounts.get("foodpanda", {}).get("name", "應收帳款-Foodpanda")

    touding_code = accounts.get("touding", {}).get("code", "1130-05")
    touding_name = accounts.get("touding", {}).get("name", "應收帳款-偷訂")

    # 1. 下載並讀取玉山銀行明細
    y, m = target_month.split("-")
    roc_tab = f"{int(y)-1911}.{m}"
    print(f"📥 下載玉山銀行明細表，讀取分頁: 【{roc_tab}】...")

    r = requests.get(f"https://www.googleapis.com/drive/v3/files/{bank_file_id}?alt=media", headers=headers)
    if r.status_code != 200:
        print(f"❌ 無法下載玉山銀行檔案 (HTTP {r.status_code})")
        return []

    wb = openpyxl.load_workbook(io.BytesIO(r.content))
    if roc_tab not in wb.sheetnames:
        print(f"⚠️ 玉山銀行檔案中找不到分頁【{roc_tab}】，可用分頁: {wb.sheetnames}")
        return []

    ws_bank = wb[roc_tab]
    bank_deposits = []
    for row in ws_bank.iter_rows(min_row=2, values_only=True):
        date_raw = row[2]
        desc = str(row[4]).strip() if row[4] else ""
        dep_raw = row[6]
        remark = str(row[8]).strip() if row[8] is not None else ""
        acct = str(row[9]).strip() if len(row) > 9 and row[9] is not None else ""
        if dep_raw is not None:
            try:
                amt = float(str(dep_raw).replace(",", ""))
                if amt > 0:
                    d_str = date_raw.strftime("%Y-%m-%d") if hasattr(date_raw, "strftime") else str(date_raw)[:10]
                    bank_deposits.append({
                        "date": d_str,
                        "desc": desc,
                        "amount": amt,
                        "remark": remark,
                        "acct": acct,
                        "used": False
                    })
            except:
                pass

    print(f"📊 玉山銀行【{roc_tab}】共讀取到 {len(bank_deposits)} 筆入帳明細。")

    # 2. 讀取門市營運交班總表
    store_sheet_id = store_cfg.get("spreadsheet_id")
    sh = gc.open_by_key(store_sheet_id)
    ws_store = [w for w in sh.worksheets() if w.title == target_month][0]
    store_rows = ws_store.get_all_values()
    store_headers = store_rows[1]

    cash_idx = store_headers.index("現金") if "現金" in store_headers else None
    deposit_time_idx = store_headers.index("存款時間") + 1 if "存款時間" in store_headers else None
    deposit_amt_idx = store_headers.index("存款金額") + 1 if "存款金額" in store_headers else None
    deposit_voucher_idx = store_headers.index("存款傳票號碼") + 1 if "存款傳票號碼" in store_headers else None

    # 3. 勾稽現金存款
    matched_cash = []
    unmatched_store_records = []
    reconciled_vouchers = []
    v_prefix = store_cfg.get("voucher_prefix", "F")
    full_cfg = load_config()
    use_auto_no = full_cfg.get("voucher_no_mode") == "@"
    
    # 傳票號碼流水號 (若為 @ 模式則輸出 @，否則從 0051 開始流水號)
    date_seq_map = {}
    def get_next_voucher_no(roc_compact_str):
        if use_auto_no:
            return "@"
        if roc_compact_str not in date_seq_map:
            date_seq_map[roc_compact_str] = 51
        seq_val = date_seq_map[roc_compact_str]
        date_seq_map[roc_compact_str] += 1
        return f"{v_prefix}{roc_compact_str}{seq_val:04d}"

    print("\n🔍 【現金存款自動勾稽與對帳】：")
    cells_to_write = []

    # 建立門市每日實際現金營收紀錄 (包含上月底跨月款項)
    daily_cash_records = []

    y_str, m_str = target_month.split("-")
    if int(m_str) == 1:
        prev_month = f"{int(y_str)-1}-12"
    else:
        prev_month = f"{y_str}-{int(m_str)-1:02d}"

    # 1. 讀取上月最後 7 天現金紀錄 (供跨月比對)
    prev_ws_list = [w for w in sh.worksheets() if w.title == prev_month]
    if prev_ws_list:
        try:
            prev_rows = prev_ws_list[0].get_all_values()
            prev_headers = prev_rows[1]
            p_cash_idx = prev_headers.index("現金") if "現金" in prev_headers else None
            if p_cash_idx:
                prev_c_list = []
                for r in prev_rows[2:]:
                    if len(r) > 1 and r[1] and any(c.isdigit() for c in r[1]) and "小計" not in r[1]:
                        d_val = r[1].strip()
                        c_val = r[p_cash_idx].replace("$", "").replace(",", "").strip() if len(r) > p_cash_idx else ""
                        try:
                            c_amt = float(c_val)
                            if c_amt > 0:
                                prev_c_list.append({
                                    "row_idx": None,
                                    "date": d_val,
                                    "cash": c_amt,
                                    "used": False,
                                    "is_cross_month": True
                                })
                        except:
                            pass
                daily_cash_records.extend(prev_c_list[-7:])
        except Exception as pe:
            print(f"⚠️ 讀取上月 ({prev_month}) 試算表失敗: {pe}")

    # 2. 讀取當月每日現金紀錄
    for r_idx, r in enumerate(store_rows[2:], start=3):
        d_val = r[1].strip()
        if d_val and any(c.isdigit() for c in d_val) and "小計" not in d_val:
            c_val = r[cash_idx].replace("$", "").replace(",", "").strip() if cash_idx and len(r) > cash_idx else ""
            try:
                c_amt = float(c_val)
                daily_cash_records.append({
                    "row_idx": r_idx,
                    "date": d_val,
                    "cash": c_amt,
                    "used": False,
                    "is_cross_month": False
                })
            except:
                pass

    # 3. 現金存款自動勾稽 (支援跨月單日、跨月多日累計、當月多日連存與短溢調節)
    cso_cfg = full_cfg.get("cash_short_over", {})
    cso_code = cso_cfg.get("code", "7112")
    cso_name = cso_cfg.get("name", "現金短溢")
    tolerance = float(cso_cfg.get("tolerance", 500))

    # 遍歷玉山銀行所有的現金存入項目
    for b in bank_deposits:
        if b["used"] or "現金" not in b["desc"]:
            continue
        
        b_amt = b["amount"]
        try:
            b_dt = datetime.strptime(b["date"], "%Y-%m-%d").date()
        except:
            b_dt = None

        match_found = False
        matched_chunk = []
        diff = 0.0

        # Pass 1: 單日金額完全相符 (容差 <= 1.0，存款日必須在營收日 0~7 天內)
        for drec in daily_cash_records:
            if not drec["used"] and abs(drec["cash"] - b_amt) <= 1.0:
                try:
                    s_dt = datetime.strptime(drec["date"], "%Y-%m-%d").date()
                    if b_dt and 0 <= (b_dt - s_dt).days <= 7:
                        matched_chunk = [drec]
                        diff = round(b_amt - drec["cash"], 2)
                        match_found = True
                        break
                    elif not b_dt:
                        matched_chunk = [drec]
                        diff = round(b_amt - drec["cash"], 2)
                        match_found = True
                        break
                except:
                    pass

        # Pass 2: 連續多日累計「完全相符」 (容差 <= 1.0，優先配對無尾差之多日累計，例如 09/30+10/01 剛好 $12,260)
        if not match_found:
            for w in range(2, 8):
                for i in range(len(daily_cash_records) - w + 1):
                    chunk = daily_cash_records[i:i+w]
                    if any(c["used"] for c in chunk):
                        continue
                    try:
                        c_end_dt = datetime.strptime(chunk[-1]["date"], "%Y-%m-%d").date()
                        if b_dt and not (0 <= (b_dt - c_end_dt).days <= 7):
                            continue
                    except:
                        pass
                    
                    c_sum = sum(c["cash"] for c in chunk)
                    cur_diff = round(b_amt - c_sum, 2)
                    if abs(cur_diff) <= 1.0:
                        matched_chunk = chunk
                        diff = cur_diff
                        match_found = True
                        break
                if match_found:
                    break

        # Pass 3: 連續多日累計在容許度範圍內 (例如多日連存差幾十塊短溢調節)
        if not match_found:
            for w in range(2, 8):
                for i in range(len(daily_cash_records) - w + 1):
                    chunk = daily_cash_records[i:i+w]
                    if any(c["used"] for c in chunk):
                        continue
                    try:
                        c_end_dt = datetime.strptime(chunk[-1]["date"], "%Y-%m-%d").date()
                        if b_dt and not (0 <= (b_dt - c_end_dt).days <= 7):
                            continue
                    except:
                        pass
                    
                    c_sum = sum(c["cash"] for c in chunk)
                    cur_diff = round(b_amt - c_sum, 2)
                    if abs(cur_diff) <= tolerance:
                        matched_chunk = chunk
                        diff = cur_diff
                        match_found = True
                        break
                if match_found:
                    break

        # Pass 4: 單日金額在容許度範圍內 (單日短溢調節)
        if not match_found:
            for drec in daily_cash_records:
                if not drec["used"]:
                    try:
                        s_dt = datetime.strptime(drec["date"], "%Y-%m-%d").date()
                        if b_dt and 0 <= (b_dt - s_dt).days <= 5:
                            cur_diff = round(b_amt - drec["cash"], 2)
                            if abs(cur_diff) <= tolerance:
                                matched_chunk = [drec]
                                diff = cur_diff
                                match_found = True
                                break
                    except:
                        pass

        # 執行配對與傳票生成
        if match_found and matched_chunk:
            b["used"] = True
            for c in matched_chunk:
                c["used"] = True
                c["deposit_date"] = b["date"]
                c["deposit_amount"] = c["cash"]
                c["diff"] = diff if len(matched_chunk) == 1 else 0.0

            chunk_sum = sum(c["cash"] for c in matched_chunk)
            
            # 格式化所屬款項日描述 (清晰識別跨月款或多日款)
            if len(matched_chunk) == 1:
                c_single = matched_chunk[0]
                if c_single.get("is_cross_month"):
                    period_str = f"{c_single['date']} (上月跨月款)"
                else:
                    period_str = c_single["date"]
            else:
                has_cross = any(c.get("is_cross_month") for c in matched_chunk)
                if has_cross:
                    period_str = f"{matched_chunk[0]['date']}~{matched_chunk[-1]['date']} (跨月多日累計)"
                else:
                    period_str = f"{matched_chunk[0]['date']}~{matched_chunk[-1]['date']} 多日累計"

            for c in matched_chunk:
                c["period"] = period_str

            diff_str = f" [調節短溢: ${diff:+.0f}]" if abs(diff) > 0.5 else ""
            print(f"   ✅ 門市【{period_str}】現金款項: ${chunk_sum:>8,.0f} ➔ 玉山 {b['date']} 存入: ${b['amount']:>8,.0f} (備註: {b['remark']}){diff_str}")

            roc_date_slash, roc_compact = to_roc_date(b["date"])
            v_no = get_next_voucher_no(roc_compact)

            # 借方: 銀行存款-玉山 1102 (實際入帳金額)
            reconciled_vouchers.append({
                "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                "總摘要": f"門市營業現金存入銀行_{store_name}_{period_str}",
                "項次": "0001", "會計項目": bank_acct_code, "項目名稱": bank_acct_name, "借貸": "借",
                "對象類別": "內部", "對象": store_name, "摘要": f"營業現金存入玉山銀行 (所屬款項日: {period_str})",
                "本幣金額": round(b["amount"]), "部門": store_cfg.get("department", store_name), "專案": ""
            })

            seq = 2
            # 現金短少: 借方增列 現金短溢
            if diff < -0.5:
                reconciled_vouchers.append({
                    "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                    "總摘要": f"門市營業現金存入銀行_{store_name}_{period_str}",
                    "項次": f"{seq:04d}", "會計項目": cso_code, "項目名稱": cso_name, "借貸": "借",
                    "對象類別": "內部", "對象": store_name, "摘要": f"門市營業現金短少調節 (差額 ${abs(diff):,.0f})",
                    "本幣金額": round(abs(diff)), "部門": store_cfg.get("department", store_name), "專案": ""
                })
                seq += 1
            # 現金溢收: 貸方增列 現金短溢
            elif diff > 0.5:
                reconciled_vouchers.append({
                    "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                    "總摘要": f"門市營業現金存入銀行_{store_name}_{period_str}",
                    "項次": f"{seq:04d}", "會計項目": cso_code, "項目名稱": cso_name, "借貸": "貸",
                    "對象類別": "內部", "對象": store_name, "摘要": f"門市營業現金溢收調節 (差額 ${diff:,.0f})",
                    "本幣金額": round(diff), "部門": store_cfg.get("department", store_name), "專案": ""
                })
                seq += 1

            # 貸方: 沖銷門市現金應收 (chunk_sum)
            reconciled_vouchers.append({
                "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                "總摘要": f"門市營業現金存入銀行_{store_name}_{period_str}",
                "項次": f"{seq:04d}", "會計項目": cash_code, "項目名稱": cash_name, "借貸": "貸",
                "對象類別": "客戶", "對象": "一般顧客", "摘要": f"沖銷 {period_str} 營業現金應收 (玉山 {b['date']} 存入)",
                "本幣金額": round(chunk_sum), "部門": store_cfg.get("department", store_name), "專案": ""
            })

            # 回填 Google 試算表 (針對當月列)
            # 回填 Google 試算表（依指示：僅回填傳票號碼，其餘欄位如存款時間、存款金額等一律不動）
            if not preview:
                for c in matched_chunk:
                    if c.get("row_idx") and deposit_voucher_idx:
                        r_idx = c["row_idx"]
                        cells_to_write.append(gspread.Cell(row=r_idx, col=deposit_voucher_idx, value=v_no))

    # 4. 勾稽第三方金流撥款 (LINE Pay, LINE Pay Money 387, Uber Eats, Foodpanda, 偷訂)
    print("\n🔍 【第三方金流 (LINE Pay / 外送 / 偷訂) 入帳勾稽】：")
    for b in bank_deposits:
        if b["used"]:
            continue
        roc_date_slash, roc_compact = to_roc_date(b["date"])

        # Case 1: LINE Pay (連加網路信用卡/第三方金流撥款 - 統編 99024139 或 連加網路)
        if "99024139" in b["remark"] or "連加網路" in b["desc"]:
            b["used"] = True
            v_no = get_next_voucher_no(roc_compact)
            net_amt = b["amount"]
            gross_amt = round(net_amt / 0.9685)  # 扣約 3.15% 總稅費推算
            fee_total = gross_amt - net_amt
            fee_net = round(fee_total / 1.05)
            fee_tax = fee_total - fee_net

            # 計算所屬營業日 (LINE Pay 信用卡通常為 T-2 營業日撥款)
            try:
                b_dt = datetime.strptime(b["date"], "%Y-%m-%d")
                days_back = 4 if b_dt.weekday() == 0 else (3 if b_dt.weekday() == 1 else 2)
                s_dt = b_dt - timedelta(days=days_back)
                sales_period_str = f"約{s_dt.strftime('%m/%d')}營業款"
            except:
                sales_period_str = "所屬營業日"

            print(f"   💚 【LINE Pay 撥款】入帳日: {b['date']} ({sales_period_str}) | 實收: ${net_amt:>8,.0f} | 預估手續費: ${fee_net:>5,.0f} (稅: ${fee_tax:>3,.0f})")

            reconciled_vouchers.append({
                "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                "總摘要": f"連加網路(LINE Pay)撥款入帳_{store_name}_{sales_period_str}",
                "項次": "0001", "會計項目": bank_acct_code, "項目名稱": bank_acct_name, "借貸": "借",
                "對象類別": "廠商", "對象": "連加網路", "摘要": f"LINE Pay 撥款入帳淨額 (款項日: {sales_period_str})",
                "本幣金額": round(net_amt), "部門": store_cfg.get("department", store_name), "專案": ""
            })
            reconciled_vouchers.append({
                "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                "總摘要": f"連加網路(LINE Pay)撥款入帳_{store_name}_{sales_period_str}",
                "項次": "0002", "會計項目": fee_acct_code, "項目名稱": fee_acct_name, "借貸": "借",
                "對象類別": "廠商", "對象": "連加網路", "摘要": f"LINE Pay 交易手續費",
                "本幣金額": round(fee_net), "部門": store_cfg.get("department", store_name), "專案": ""
            })
            reconciled_vouchers.append({
                "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                "總摘要": f"連加網路(LINE Pay)撥款入帳_{store_name}_{sales_period_str}",
                "項次": "0003", "會計項目": fee_tax_code, "項目名稱": fee_tax_name, "借貸": "借",
                "對象類別": "廠商", "對象": "連加網路", "摘要": f"LINE Pay 手續費進項稅額",
                "本幣金額": round(fee_tax), "部門": store_cfg.get("department", store_name), "專案": ""
            })
            reconciled_vouchers.append({
                "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                "總摘要": f"連加網路(LINE Pay)撥款入帳_{store_name}_{sales_period_str}",
                "項次": "0004", "會計項目": linepay_code, "項目名稱": linepay_name, "借貸": "貸",
                "對象類別": "客戶", "對象": "一般顧客", "摘要": f"沖銷 {sales_period_str} LINE Pay 應收帳款",
                "本幣金額": round(gross_amt), "部門": store_cfg.get("department", store_name), "專案": ""
            })

        # Case 2: 連加電子支付 (387 / LINE Pay Money 電支提領入帳)
        elif b["acct"].startswith("387") or "387" in b["remark"] or "連加電子" in b["desc"]:
            b["used"] = True
            v_no = get_next_voucher_no(roc_compact)
            try:
                b_dt = datetime.strptime(b["date"], "%Y-%m-%d")
                s_dt = b_dt - timedelta(days=1)
                sales_period_str = f"約{s_dt.strftime('%m/%d')}電支款"
            except:
                sales_period_str = "所屬電支款"
            print(f"   🟢 【LINE Pay 電支】入帳日: {b['date']} ({sales_period_str}) | 提領入帳: ${b['amount']:>8,.0f}")
            reconciled_vouchers.append({
                "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                "總摘要": f"連加電子支付(LINE Pay)提領入帳_{store_name}_{sales_period_str}",
                "項次": "0001", "會計項目": bank_acct_code, "項目名稱": bank_acct_name, "借貸": "借",
                "對象類別": "廠商", "對象": "連加電子支付", "摘要": f"LINE Pay 電子支付提領 (款項日: {sales_period_str})",
                "本幣金額": round(b["amount"]), "部門": store_cfg.get("department", store_name), "專案": ""
            })
            reconciled_vouchers.append({
                "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                "總摘要": f"連加電子支付(LINE Pay)提領入帳_{store_name}_{sales_period_str}",
                "項次": "0002", "會計項目": linepay_code, "項目名稱": linepay_name, "借貸": "貸",
                "對象類別": "客戶", "對象": "一般顧客", "摘要": f"沖銷 {sales_period_str} LINE Pay 應收帳款",
                "本幣金額": round(b["amount"]), "部門": store_cfg.get("department", store_name), "專案": ""
            })

        # Case 3: Uber Eats (統編 83118125 或 優步)
        elif "83118125" in b["remark"] or "優步" in b["desc"]:
            b["used"] = True
            v_no = get_next_voucher_no(roc_compact)
            try:
                b_dt = datetime.strptime(b["date"], "%Y-%m-%d")
                s_start = b_dt - timedelta(days=14)
                s_end = b_dt - timedelta(days=8)
                sales_period_str = f"{s_start.strftime('%m/%d')}~{s_end.strftime('%m/%d')}外送款"
            except:
                sales_period_str = "所屬外送款"
            print(f"   🚗 【Uber Eats 撥款】入帳日: {b['date']} ({sales_period_str}) | 實收淨額: ${b['amount']:>8,.0f}")
            reconciled_vouchers.append({
                "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                "總摘要": f"Uber Eats 外送撥款入帳_{store_name}_{sales_period_str}",
                "項次": "0001", "會計項目": bank_acct_code, "項目名稱": bank_acct_name, "借貸": "借",
                "對象類別": "廠商", "對象": "優步福爾摩沙", "摘要": f"Uber Eats 外送撥款 (款項日: {sales_period_str})",
                "本幣金額": round(b["amount"]), "部門": store_cfg.get("department", store_name), "專案": ""
            })
            reconciled_vouchers.append({
                "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                "總摘要": f"Uber Eats 外送撥款入帳_{store_name}_{sales_period_str}",
                "項次": "0002", "會計項目": ubereats_code, "項目名稱": ubereats_name, "借貸": "貸",
                "對象類別": "客戶", "對象": "一般顧客", "摘要": f"沖銷 {sales_period_str} Uber Eats 應收帳款",
                "本幣金額": round(b["amount"]), "部門": store_cfg.get("department", store_name), "專案": ""
            })

        # Case 4: Foodpanda (富胖達)
        elif "富胖達" in b["remark"] or "富胖達" in b["desc"]:
            b["used"] = True
            v_no = get_next_voucher_no(roc_compact)
            try:
                b_dt = datetime.strptime(b["date"], "%Y-%m-%d")
                sales_period_str = "上月下半月(16~30日)外送款" if b_dt.day <= 15 else "本月上半月(01~15日)外送款"
            except:
                sales_period_str = "所屬外送款"
            print(f"   🐼 【Foodpanda 撥款】入帳日: {b['date']} ({sales_period_str}) | 實收淨額: ${b['amount']:>8,.0f}")
            reconciled_vouchers.append({
                "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                "總摘要": f"Foodpanda 外送撥款入帳_{store_name}_{sales_period_str}",
                "項次": "0001", "會計項目": bank_acct_code, "項目名稱": bank_acct_name, "借貸": "借",
                "對象類別": "廠商", "對象": "富胖達", "摘要": f"Foodpanda 外送撥款 (款項日: {sales_period_str})",
                "本幣金額": round(b["amount"]), "部門": store_cfg.get("department", store_name), "專案": ""
            })
            reconciled_vouchers.append({
                "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                "總摘要": f"Foodpanda 外送撥款入帳_{store_name}_{sales_period_str}",
                "項次": "0002", "會計項目": foodpanda_code, "項目名稱": foodpanda_name, "借貸": "貸",
                "對象類別": "客戶", "對象": "一般顧客", "摘要": f"沖銷 {sales_period_str} Foodpanda 應收帳款",
                "本幣金額": round(b["amount"]), "部門": store_cfg.get("department", store_name), "專案": ""
            })

        # Case 5: 偷訂餐飲科技 (統編 93633536 或 偷訂)
        elif "93633536" in b["remark"] or "偷訂" in b["desc"] or "偷訂" in b["remark"]:
            b["used"] = True
            v_no = get_next_voucher_no(roc_compact)
            print(f"   🛒 【偷訂科技 撥款】入帳日: {b['date']} | 實收淨額: ${b['amount']:>8,.0f}")
            reconciled_vouchers.append({
                "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                "總摘要": f"偷訂餐飲科技撥款入帳_{store_name}",
                "項次": "0001", "會計項目": bank_acct_code, "項目名稱": bank_acct_name, "借貸": "借",
                "對象類別": "廠商", "對象": "偷訂餐飲科技", "摘要": f"偷訂平台撥款入帳",
                "本幣金額": round(b["amount"]), "部門": store_cfg.get("department", store_name), "專案": ""
            })
            reconciled_vouchers.append({
                "傳票號碼": v_no, "傳票種類": "轉帳", "傳票日期": roc_date_slash, "附件數": 0,
                "總摘要": f"偷訂餐飲科技撥款入帳_{store_name}",
                "項次": "0002", "會計項目": touding_code, "項目名稱": touding_name, "借貸": "貸",
                "對象類別": "客戶", "對象": "一般顧客", "摘要": f"沖銷 偷訂 應收帳款",
                "本幣金額": round(b["amount"]), "部門": store_cfg.get("department", store_name), "專案": ""
            })

    # 5. 回填試算表（僅回填存款傳票號碼）
    if cells_to_write and not preview:
        try:
            ws_store.update_cells(cells_to_write)
            print(f"\n✍️ 成功自動回填 {len(cells_to_write)} 筆【存款傳票號碼】至【{store_name}】營運表！（其餘欄位完全不更動）")
        except Exception as e:
            print(f"⚠️ 回填失敗: {e}")

    return reconciled_vouchers, daily_cash_records

def main():
    parser = argparse.ArgumentParser(description="玉山銀行明細智慧對帳與沖銷傳票產生")
    parser.add_argument("--month", type=str, default="2026-10", help="指定對帳月份（預設 2026-10）")
    parser.add_argument("--preview", action="store_true", help="僅預覽對帳結果，不產生傳票檔也不回填")
    args = parser.parse_args()

    full_config = load_config()
    stores = full_config.get("stores", {})

    all_vouchers = []
    all_store_cash = {}
    for s_name, s_cfg in stores.items():
        res = reconcile_store_bank(s_name, s_cfg, target_month=args.month, preview=args.preview)
        if isinstance(res, tuple):
            vouchers, daily_cash = res
        else:
            vouchers, daily_cash = res, []
        if vouchers:
            all_vouchers.extend(vouchers)
        all_store_cash[s_name] = daily_cash

    # 僅產出單一【全部門市合併 Excel】
    if all_vouchers and not args.preview:
        import openpyxl
        os.makedirs(OUTPUT_DIR, exist_ok=True)
        compact_month = args.month.replace("-", "")
        out_all_excel = os.path.join(OUTPUT_DIR, f"文中傳票維護匯入_全部合併_玉山銀行沖銷_{compact_month}.xlsx")

        # Excel
        wb_all = openpyxl.Workbook()
        ws_all = wb_all.active
        ws_all.title = "銀行沖銷傳票_全部合併"
        ws_all.append(OFFICIAL_HEADER_1)
        ws_all.append(OFFICIAL_HEADER_2)
        for rv in all_vouchers:
            ws_all.append(format_official_42_row(rv))
        wb_all.save(out_all_excel)

        print("\n" + "="*70)
        print(f"🌟 全部門市合併銀行沖銷傳票已產出（共 {len(all_vouchers)} 筆分錄）：")
        print(f"   💾 總合併 Excel: {os.path.abspath(out_all_excel)}")
        print("="*70)

        # 自動更新並顯示本月對帳差額與應收餘額表
        try:
            import reconciliation_summary
            summary_data = reconciliation_summary.calculate_monthly_summary(target_month=args.month, store_cash_data=all_store_cash)
            if summary_data:
                reconciliation_summary.print_and_export_summary(summary_data, target_month=args.month)
        except Exception as se:
            print(f"⚠️ 更新對帳差額表時發生錯誤: {se}")


if __name__ == "__main__":
    main()
