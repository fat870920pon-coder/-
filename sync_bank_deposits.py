#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
sync_bank_deposits.py
---------------------
玉山銀行存摺明細檔同步讀取程式（非登入網銀，係從 Google Drive 下載指定銀行存摺 Excel 明細檔）。
涵蓋商戶：
- 月宮（月下餐飲股份有限公司）：File ID 1-6Tff22lTh5zJ8JPlgjeTHOQrQp_tmq6
- 隨意小吃（頑全公關顧問股份有限公司）：File ID 1jI-5x-6jOViRrWkVfOOA3izKrHUNnBaW

審計合規要求：
- 若下載失敗或連線異常，嚴禁將存款當成零或隱匿失敗。
- 失敗時保留原有成功快取，不得將離線重算或快取當成同步成功。
- 成功時更新 cached_esun_bank_deposits.json 並計算資料涵蓋日期區間。
"""

import os
import sys
import io
import json
import time
from datetime import datetime
import requests

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
CRED_PATH = os.path.join(BASE_DIR, "credentials.json")
if not os.path.exists(CRED_PATH):
    ALT_CRED = os.path.expanduser("~/Desktop/腳本資料夾/月下/文中會計傳票產生器/credentials.json")
    if os.path.exists(ALT_CRED):
        CRED_PATH = ALT_CRED

CACHED_BANK_PATH = os.path.join(BASE_DIR, "cached_esun_bank_deposits.json")
ALT_CACHED_BANK_PATH = os.path.expanduser("~/Desktop/腳本資料夾/月下/文中會計傳票產生器/cached_esun_bank_deposits.json")

BANK_FILE_IDS = {
    "suiyi": "1jI-5x-6jOViRrWkVfOOA3izKrHUNnBaW",
    "yuegong": "1-6Tff22lTh5zJ8JPlgjeTHOQrQp_tmq6"
}

def get_drive_credentials():
    if not os.path.exists(CRED_PATH):
        raise FileNotFoundError(f"找不到 Google 憑證檔案：{CRED_PATH}")
    try:
        from google.oauth2.service_account import Credentials
        from google.auth.transport.requests import Request
        scopes = ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive"]
        creds = Credentials.from_service_account_file(CRED_PATH, scopes=scopes)
        creds.refresh(Request())
        return creds.token
    except Exception as e:
        raise RuntimeError(f"取得 Google Drive 授權失敗：{e}")

def sync_bank_deposits():
    now_str = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    print(f"[{datetime.now().strftime('%H:%M:%S')}] 🏦 開始自 Google Drive 讀取指定玉山銀行明細檔...")

    try:
        import openpyxl
    except ImportError:
        print("❌ 執行環境缺少 openpyxl，無法解析 Excel 明細檔")
        return {"ok": False, "error": "缺少 openpyxl 套件"}

    token = get_drive_credentials()
    headers = {"Authorization": f"Bearer {token}"}
    results = {}
    fetch_errors = []
    all_dates = []

    for store_key, file_id in BANK_FILE_IDS.items():
        results[store_key] = []
        try:
            print(f"  📥 下載商戶 [{store_key}] 玉山存摺明細檔 (File ID: {file_id})...")
            r = requests.get(
                f"https://www.googleapis.com/drive/v3/files/{file_id}?alt=media",
                headers=headers,
                timeout=45
            )
            if r.status_code == 200:
                wb = openpyxl.load_workbook(io.BytesIO(r.content), data_only=True)
                for sname in wb.sheetnames:
                    ws = wb[sname]
                    for row in list(ws.iter_rows(values_only=True))[1:]:
                        if not any(row):
                            continue
                        date_val = row[2]
                        if not date_val:
                            continue
                        d_str = str(date_val)[:10].replace("/", "-")
                        dep_raw = row[6]
                        desc = str(row[4]).strip() if row[4] else ""
                        rem = str(row[8]).strip() if len(row) > 8 and row[8] is not None else ""
                        try:
                            val = float(str(dep_raw).replace(",", ""))
                            if val > 0:
                                results[store_key].append({
                                    "sheet": sname,
                                    "date": d_str,
                                    "amount": val,
                                    "desc": desc,
                                    "remark": rem,
                                    "raw_row": [str(c) if c is not None else "" for c in row[:10]]
                                })
                                all_dates.append(d_str)
                        except Exception:
                            pass
                print(f"  ✅ [{store_key}] 解析成功：共 {len(results[store_key])} 筆存入流水")
            else:
                err = f"商戶 [{store_key}] 下載失敗 HTTP {r.status_code}"
                print(f"  ❌ {err}")
                fetch_errors.append(err)
        except Exception as e:
            err = f"商戶 [{store_key}] 下載解析異常: {str(e)}"
            print(f"  ❌ {err}")
            fetch_errors.append(err)

    if fetch_errors:
        err_summary = "; ".join(fetch_errors)
        print(f"⚠️ 銀行明細檔更新未完全成功：{err_summary}。保留既有快取。")
        return {"ok": False, "error": err_summary}

    # 計算資料涵蓋日期
    date_coverage = "未知"
    valid_dates = [d for d in all_dates if len(d) == 10 and d.startswith("20")]
    if valid_dates:
        valid_dates.sort()
        date_coverage = f"{valid_dates[0]} ~ {valid_dates[-1]}"

    # 寫入快取檔案
    with open(CACHED_BANK_PATH, "w", encoding="utf-8") as f:
        json.dump(results, f, ensure_ascii=False, indent=2)

    try:
        with open(ALT_CACHED_BANK_PATH, "w", encoding="utf-8") as f:
            json.dump(results, f, ensure_ascii=False, indent=2)
    except Exception:
        pass

    total_count = sum(len(v) for v in results.values())
    print(f"[{datetime.now().strftime('%H:%M:%S')}] ✅ 玉山銀行明細檔更新完成！共 {total_count} 筆，資料涵蓋：{date_coverage}")
    return {
        "ok": True,
        "total_records": total_count,
        "data_coverage_range": date_coverage,
        "updated_at": now_str
    }

if __name__ == "__main__":
    res = sync_bank_deposits()
    if not res["ok"]:
        sys.exit(1)
