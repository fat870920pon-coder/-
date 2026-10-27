import os
import sys
import json
import re
from datetime import datetime, date, timedelta
import openpyxl
import shutil

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
OUTPUT_DIR = os.path.join(BASE_DIR, "output")
STATE_FILE = os.path.join(BASE_DIR, "reconciliation_balance.json")

def load_month_payload(target_month, state_data):
    m_clean = target_month.replace("-", "")
    month_data = state_data.get(target_month, {})
    updated_at = month_data.get("updated_at", datetime.now().strftime("%Y-%m-%d %H:%M:%S"))
    stores_data = month_data.get("stores", {})

    bank_excel = os.path.join(OUTPUT_DIR, f"本月對帳差額與應收餘額表_{m_clean}.xlsx")
    cleared_items = []
    if os.path.exists(bank_excel):
        try:
            wb = openpyxl.load_workbook(bank_excel, data_only=True)
            if "玉山單筆沖銷對照" in wb.sheetnames:
                ws2 = wb["玉山單筆沖銷對照"]
                for r in list(ws2.iter_rows(values_only=True))[1:]:
                    if r[0] is not None:
                        alloc_val = str(r[7] or "") if len(r) >= 9 else ""
                        rem_val = str(r[8] or "") if len(r) >= 9 else str(r[7] or "")
                        is_cross = bool("跨月" in alloc_val or "上月" in alloc_val or "拆分" in alloc_val)
                        cleared_items.append({
                            "idx": r[0],
                            "store": r[1],
                            "bank_date": str(r[2]),
                            "category": str(r[3]),
                            "bank_amt": float(r[4] or 0),
                            "cleared_amt": float(r[5] or 0),
                            "period": str(r[6]),
                            "revenue_alloc": alloc_val,
                            "is_cross_month": is_cross,
                            "remark": rem_val
                        })
        except Exception:
            pass

    total_pos = 0.0
    total_bank = 0.0
    total_inv = 0.0
    total_voucher_fee = 0.0
    total_est_fee = 0.0
    total_est_vat = 0.0
    total_pending_cash = 0.0

    total_unpaid_cash = 0.0
    total_declared_pending = 0.0
    all_pending_alerts = []
    all_invoice_tracking = []
    for skey, sinfo in stores_data.items():
        if skey.startswith("_"):
            continue
        channels = sinfo.get("channels", {})
        for ccode, cdata in channels.items():
            total_pos += cdata.get("pos_sales", 0.0)
            total_bank += cdata.get("bank_cleared", 0.0)
            total_inv += cdata.get("invoice_cleared", 0.0)
            total_voucher_fee += cdata.get("voucher_fee_tax", 0.0)
            total_est_fee += cdata.get("est_total_fee", 0.0)
            if ccode == "1130-01":
                diff = cdata.get("diff", 0.0)
                if diff > 0:
                    total_pending_cash += diff

        for d in sinfo.get("daily_cash", []):
            st = d.get("audit_status", "")
            amt = d.get("cash", 0.0)
            if st == "UNPAID":
                total_unpaid_cash += amt
            elif st == "DECLARED_PENDING_BANK":
                total_declared_pending += amt
        
        alerts = sinfo.get("pending_alerts", [])
        for a in alerts:
            all_pending_alerts.append(a)

        for it in sinfo.get("invoice_tracking", []):
            total_est_vat += it.get("est_vat_5pct", 0.0)
            all_invoice_tracking.append(it)

    all_pending_alerts.sort(key=lambda x: (0 if x.get("audit_status") == "UNPAID" else 1, -x.get("delay_days", 0)))
    bridge_data = month_data.get("_bridge", {})

    return {
        "target_month": target_month,
        "updated_at": updated_at,
        "kpi": {
            "total_pos": total_pos,
            "total_bank": total_bank,
            "total_inv": total_inv,
            "total_voucher_fee": total_voucher_fee,
            "total_est_fee": total_est_fee,
            "total_est_vat": total_est_vat,
            "total_pending_cash": total_pending_cash,
            "total_unpaid_cash": total_unpaid_cash,
            "total_declared_pending": total_declared_pending,
            "alert_count": len(all_pending_alerts)
        },
        "bridge": bridge_data,
        "stores": stores_data,
        "alerts": all_pending_alerts,
        "invoice_tracking": all_invoice_tracking,
        "cleared_items": cleared_items
    }

def generate_html_dashboard(target_month="2026-10", sync_cloud=False):
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    m_clean = target_month.replace("-", "")

    state_data = {}
    if os.path.exists(STATE_FILE):
        with open(STATE_FILE, "r", encoding="utf-8") as f:
            state_data = json.load(f)

    all_months = sorted([k for k in state_data.keys() if re.match(r"^\d{4}-\d{2}$", str(k))], reverse=True)
    if target_month not in all_months:
        all_months.insert(0, target_month)

    multi_months_data = {}
    for m in all_months:
        multi_months_data[m] = load_month_payload(m, state_data)

    current_data = multi_months_data.get(target_month, multi_months_data[all_months[0]])

    json_payload_str = json.dumps(multi_months_data, ensure_ascii=False)
    for json_path in [
        os.path.join(BASE_DIR, "dashboard_data.json"),
        os.path.join(BASE_DIR, "docs", "dashboard_data.json"),
        os.path.join(OUTPUT_DIR, "dashboard_data.json")
    ]:
        os.makedirs(os.path.dirname(json_path), exist_ok=True)
        with open(json_path, "w", encoding="utf-8") as jf:
            jf.write(json_payload_str)

    month_options_html = ""
    for m in all_months:
        y, mo = m.split("-")
        selected_attr = "selected" if m == target_month else ""
        month_options_html += f'<option value="{m}" {selected_attr}>{y} 年 {mo} 月</option>'

    html_content = f"""<!DOCTYPE html>
<html lang="zh-TW">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <meta http-equiv="Cache-Control" content="no-cache, no-store, must-revalidate">
  <meta http-equiv="Pragma" content="no-cache">
  <meta http-equiv="Expires" content="0">
  <title>餐飲集團 - 全部門市 · 財務對帳看板</title>
  <script src="https://cdn.tailwindcss.com"></script>
  <link rel="stylesheet" href="https://cdnjs.cloudflare.com/ajax/libs/font-awesome/6.4.0/css/all.min.css">
  <style>
    body {{
      font-family: -apple-system, BlinkMacSystemFont, "PingFang TC", "Noto Sans TC", "Microsoft JhengHei", sans-serif;
      background-color: #f1f5f9;
      color: #0f172a;
      -webkit-font-smoothing: antialiased;
    }}
    .card-box {{
      background: #ffffff;
      border: 1px solid #e2e8f0;
      border-radius: 16px;
      box-shadow: 0 1px 3px rgba(15, 23, 42, 0.03);
    }}
    .calendar-cell {{
      min-height: 92px;
      transition: all 0.12s ease;
    }}
    .calendar-cell:hover {{
      border-color: #6366f1;
      background-color: #f8fafc;
    }}
  </style>
</head>
<body class="p-4 sm:p-6 lg:p-8 max-w-7xl mx-auto space-y-6">

  <!-- Toast -->
  <div id="sync-toast" class="fixed top-4 right-4 z-50 hidden items-center gap-2 px-4 py-2.5 rounded-xl bg-slate-900 text-white text-xs font-medium shadow-lg">
    <i class="fa-solid fa-circle-check text-emerald-400"></i>
    <span id="sync-toast-text">已更新最新資料</span>
  </div>

  <!-- LINE Summary Copy Modal -->
  <div id="line-modal" class="fixed inset-0 z-50 hidden items-center justify-center bg-slate-900/50 p-4">
    <div class="bg-white rounded-2xl shadow-xl border border-slate-200 max-w-lg w-full overflow-hidden">
      <div class="px-5 py-3.5 bg-emerald-600 text-white flex justify-between items-center">
        <div class="font-bold text-sm flex items-center gap-2">
          <i class="fa-brands fa-line text-lg"></i>
          <span>已複製 LINE 催帳與雙層審核摘要</span>
        </div>
        <button onclick="closeLineModal()" class="text-white/80 hover:text-white cursor-pointer text-sm font-bold">✕</button>
      </div>
      <div class="p-5">
        <p class="text-xs text-slate-500 mb-2">以下文字已自動複製到剪貼簿，可直接貼上至 LINE 群組提醒店長或會計核對：</p>
        <textarea id="line-summary-textarea" rows="11" class="w-full p-3 rounded-xl bg-slate-50 border border-slate-200 font-mono text-xs text-slate-800 focus:outline-none focus:ring-2 focus:ring-emerald-500"></textarea>
        <div class="mt-4 flex justify-end gap-2">
          <button onclick="closeLineModal()" class="px-4 py-2 rounded-xl text-xs font-bold bg-slate-100 text-slate-600 hover:bg-slate-200 cursor-pointer">關閉</button>
          <button onclick="copyTextareaAgain()" class="px-4 py-2 rounded-xl text-xs font-bold bg-emerald-600 text-white hover:bg-emerald-700 cursor-pointer flex items-center gap-1.5">
            <i class="fa-regular fa-copy"></i> 再次複製文字
          </button>
        </div>
      </div>
    </div>
  </div>

  <!-- 1. Top Header & Controls -->
  <header class="card-box p-4 sm:p-5 flex flex-col lg:flex-row justify-between items-start lg:items-center gap-4">
    <div class="flex items-center gap-3.5">
      <div class="w-11 h-11 rounded-xl bg-indigo-600 text-white flex items-center justify-center shadow-sm shrink-0">
        <i class="fa-solid fa-scale-balanced text-lg"></i>
      </div>
      <div>
        <div class="flex flex-wrap items-center gap-2">
          <h1 class="text-lg sm:text-xl font-bold text-slate-900" id="header-main-title">餐飲集團-全部門市</h1>
          <span class="px-2 py-0.5 rounded-full bg-emerald-50 text-emerald-700 border border-emerald-200 text-[11px] font-semibold">
            ☁️ 雲端自動同步已啟用
          </span>
        </div>
        <p class="text-xs text-slate-500 mt-0.5">第一層：核對營運總表 U 欄存款申報 ➔ 第二層：核對玉山銀行存摺入帳</p>
      </div>
    </div>

    <div class="flex flex-wrap items-center gap-2 w-full lg:w-auto justify-between lg:justify-end">
      <!-- Store Switcher -->
      <div class="inline-flex p-1 rounded-xl bg-slate-100 border border-slate-200/80">
        <button onclick="switchStore('ALL')" id="tab-ALL" class="store-tab px-3.5 py-1.5 rounded-lg text-xs font-bold transition-all bg-white text-indigo-600 shadow-xs cursor-pointer">
          全部門市
        </button>
        <button onclick="switchStore('月宮')" id="tab-月宮" class="store-tab px-3.5 py-1.5 rounded-lg text-xs font-bold transition-all text-slate-600 hover:text-slate-900 cursor-pointer">
          月宮
        </button>
        <button onclick="switchStore('隨意小吃')" id="tab-隨意小吃" class="store-tab px-3.5 py-1.5 rounded-lg text-xs font-bold transition-all text-slate-600 hover:text-slate-900 cursor-pointer">
          隨意小吃
        </button>
      </div>

      <!-- Month, Excel Download & Refresh -->
      <div class="flex flex-wrap items-center gap-2">
        <select id="month-select" onchange="switchMonth(this.value)" class="px-3 py-1.5 bg-slate-50 border border-slate-200 rounded-xl font-bold text-slate-800 text-xs cursor-pointer focus:outline-none">
          {month_options_html}
        </select>

        <!-- Excel Download Dropdown -->
        <div class="relative">
          <button onclick="toggleDownloadMenu()" class="px-3 py-1.5 bg-emerald-50 hover:bg-emerald-100 text-emerald-700 rounded-xl text-xs font-bold transition-all flex items-center gap-1.5 border border-emerald-200 cursor-pointer">
            <i class="fa-solid fa-file-excel"></i>
            <span>下載 Excel / 傳票</span>
            <i class="fa-solid fa-chevron-down text-[10px]"></i>
          </button>
          <div id="download-menu" class="hidden absolute right-0 mt-2 w-64 bg-white rounded-xl shadow-lg border border-slate-200 py-1.5 z-40 text-xs">
            <div class="px-3 py-1 text-[10px] font-bold text-slate-400 uppercase">跨裝置直接下載當月報表</div>
            <a id="dl-summary-excel" href="./output/本月對帳差額與應收餘額表_{m_clean}.xlsx" download class="flex items-center gap-2.5 px-3.5 py-2 hover:bg-slate-50 text-slate-700 font-medium">
              <i class="fa-solid fa-table text-emerald-600 w-4"></i>
              <div>
                <div class="font-bold text-slate-800">對帳差額與雙層審核表</div>
                <div class="text-[10px] text-slate-400">含 Sheet3 門市雙層審核異常清單</div>
              </div>
            </a>
            <a id="dl-bank-voucher" href="./output/文中傳票維護匯入_全部合併_玉山銀行沖銷_{m_clean}.xlsx" download class="flex items-center gap-2.5 px-3.5 py-2 hover:bg-slate-50 text-slate-700 font-medium border-t border-slate-100">
              <i class="fa-solid fa-building-columns text-indigo-600 w-4"></i>
              <div>
                <div class="font-bold text-slate-800">文中傳票：玉山銀行沖銷檔</div>
                <div class="text-[10px] text-slate-400">可直接匯入文中會計系統</div>
              </div>
            </a>
            <a id="dl-pos-voucher" href="./output/文中傳票維護匯入_全部合併_{m_clean}.xlsx" download class="flex items-center gap-2.5 px-3.5 py-2 hover:bg-slate-50 text-slate-700 font-medium border-t border-slate-100">
              <i class="fa-solid fa-receipt text-amber-600 w-4"></i>
              <div>
                <div class="font-bold text-slate-800">文中傳票：門市營業收入檔</div>
                <div class="text-[10px] text-slate-400">每日 POS 應收立帳傳票</div>
              </div>
            </a>
          </div>
        </div>

        <div class="flex items-center gap-1.5 text-xs text-slate-500 bg-slate-50 border border-slate-200 px-3 py-1.5 rounded-xl">
          <i class="fa-regular fa-clock text-slate-400"></i>
          <span>更新時間：<b class="font-mono text-slate-700" id="top-updated-at">{current_data['updated_at']}</b></span>
        </div>

        <button onclick="refreshLiveData(true)" title="點擊直接觸發後端 Python 重新抓取 Google 試算表並更新為當下時間" class="px-3.5 py-1.5 bg-indigo-600 hover:bg-indigo-700 text-white rounded-xl text-xs font-bold transition-all flex items-center gap-1.5 shadow-sm cursor-pointer">
          <i class="fa-solid fa-rotate" id="refresh-icon"></i>
          <span id="refresh-btn-text">即時更新</span>
        </button>
      </div>
    </div>
  </header>

  <!-- 2A. All-Channel Compact Summary Bar (全店總計橫幅) -->
  <div class="card-box p-4 bg-white">
    <div class="grid grid-cols-1 sm:grid-cols-2 lg:grid-cols-4 gap-4 divide-y sm:divide-y-0 sm:divide-x divide-slate-100">
      <div class="px-2 pt-1 sm:pt-0">
        <div class="text-xs font-semibold text-slate-500 flex items-center justify-between">
          <span>① 全通路 POS 原始總額</span>
          <i class="fa-solid fa-cash-register text-slate-400"></i>
        </div>
        <div class="text-xl font-bold text-slate-900 font-mono mt-1" id="kpi-pos">$ 0</div>
        <div class="text-[11px] text-slate-400 mt-0.5">含現金、LINE Pay 及外送平台原始營收</div>
      </div>

      <div class="px-2 pt-3 sm:pt-0 sm:pl-4">
        <div class="text-xs font-semibold text-indigo-700 flex items-center justify-between">
          <span>② 扣手續費後實收總額 (應入帳)</span>
          <i class="fa-solid fa-hand-holding-dollar text-indigo-500"></i>
        </div>
        <div class="text-xl font-bold text-indigo-600 font-mono mt-1" id="kpi-net-rec">$ 0</div>
        <div class="text-[11px] text-slate-400 mt-0.5" id="kpi-fee-note">已扣除各平台手續費與抽成</div>
      </div>

      <div class="px-2 pt-3 sm:pt-0 sm:pl-4">
        <div class="text-xs font-semibold text-emerald-700 flex items-center justify-between">
          <span>③ 玉山銀行已存總額 (屬本月)</span>
          <i class="fa-solid fa-building-columns text-emerald-500"></i>
        </div>
        <div class="text-xl font-bold text-emerald-600 font-mono mt-1" id="kpi-bank">$ 0</div>
        <div class="text-[11px] text-slate-400 mt-0.5" id="kpi-bridge-note">存摺實收已扣除上月跨月款</div>
      </div>

      <div class="px-2 pt-3 sm:pt-0 sm:pl-4">
        <div class="text-xs font-semibold text-rose-600 flex items-center justify-between">
          <span>④ 推算尚未入帳總差額 (② - ③)</span>
          <span class="text-[10px] text-slate-400 font-normal">更新 <span id="header-updated-at" class="font-mono">{current_data['updated_at']}</span></span>
        </div>
        <div class="text-xl font-bold text-rose-600 font-mono mt-1" id="kpi-pending-cash">$ 0</div>
        <div class="mt-1 flex items-center gap-1.5">
          <button onclick="scrollToInvoiceSection()" class="text-[11px] px-2 py-0.5 rounded bg-blue-50 hover:bg-blue-100 text-blue-700 font-bold border border-blue-200 cursor-pointer transition-all" id="kpi-vat-sub">
            🧾 可扣抵 5% 稅額 $0 (追蹤發票)
          </button>
          <span class="hidden" id="kpi-inv">$ 0</span>
        </div>
      </div>
    </div>
  </div>

  <!-- 2B. Four Dedicated Channel Deduction Cards (四大通路「實收 vs 已存」獨立推算字卡) -->
  <div class="grid grid-cols-1 md:grid-cols-2 xl:grid-cols-4 gap-4" id="top-channel-cards">
    <!-- Rendered by JS: 💵 現金, 💚 LINE Pay, 🚗 Uber Eats, 🐼 Foodpanda -->
  </div>

  <!-- Main Modular Navigation Tabs (清晰三大功能分頁) -->
  <div class="mt-2 flex flex-wrap items-center justify-between gap-3 border-b border-slate-200 pb-2">
    <div class="inline-flex p-1.5 rounded-2xl bg-slate-100/90 border border-slate-200 gap-1.5 shadow-inner" id="main-nav-tabs">
      <button onclick="switchMainTab('cash')" id="main-tab-btn-cash" class="main-nav-tab px-4 py-2 rounded-xl text-xs sm:text-sm font-bold bg-indigo-600 text-white shadow-sm flex items-center gap-2 cursor-pointer transition-all">
        <i class="fa-solid fa-money-bill-transfer"></i>
        <span>門市現金雙層審核</span>
      </button>
      <button onclick="switchMainTab('channels')" id="main-tab-btn-channels" class="main-nav-tab px-4 py-2 rounded-xl text-xs sm:text-sm font-semibold bg-slate-100 text-slate-600 hover:bg-slate-200 flex items-center gap-2 cursor-pointer transition-all">
        <i class="fa-solid fa-chart-pie"></i>
        <span>外送與金流四步推算</span>
      </button>
      <button onclick="switchMainTab('bank')" id="main-tab-btn-bank" class="main-nav-tab px-4 py-2 rounded-xl text-xs sm:text-sm font-semibold bg-slate-100 text-slate-600 hover:bg-slate-200 flex items-center gap-2 cursor-pointer transition-all">
        <i class="fa-solid fa-building-columns"></i>
        <span>玉山銀行入帳流水</span>
      </button>
    </div>

    <!-- Quick Action / LINE Summary Button -->
    <button onclick="copyLineSummary()" class="px-3 py-2 rounded-xl bg-[#06C755] hover:bg-[#05b34c] text-white text-xs font-bold flex items-center gap-1.5 shadow-xs cursor-pointer transition-all">
      <i class="fa-brands fa-line text-sm"></i>
      <span>一鍵複製 LINE 催帳摘要</span>
    </button>
  </div>

  <!-- Tab Pane 1: 門市現金雙層審核 (含清單 / 月曆切換) -->
  <div id="tab-pane-cash" class="space-y-4">
    <section class="card-box overflow-hidden" id="alert-section">
      <div class="p-4 sm:p-5 border-b border-slate-200 flex flex-col md:flex-row justify-between items-start md:items-center gap-3 bg-slate-50/60">
        <div>
          <div class="flex items-center gap-2">
            <span class="w-2.5 h-2.5 rounded-full bg-indigo-600"></span>
            <h2 class="font-bold text-slate-900 text-base" id="alert-title-text">門市現金雙層審核明細</h2>
          </div>
          <p class="text-xs text-slate-500 mt-0.5">優先比對營運總表 U 欄「存款金額」與短溢差額，再核對玉山銀行存摺是否入帳</p>
        </div>

        <div class="flex flex-wrap items-center gap-2">
          <!-- View Toggle: Table vs Calendar -->
          <div class="inline-flex p-1 rounded-xl bg-slate-200/70 border border-slate-300/60">
            <button onclick="toggleCashView('table')" id="btn-cash-view-table" class="px-3 py-1.5 rounded-lg text-xs font-bold bg-white text-indigo-700 shadow-xs cursor-pointer transition-all">
              <i class="fa-solid fa-table-list mr-1"></i>清單視圖
            </button>
            <button onclick="toggleCashView('calendar')" id="btn-cash-view-calendar" class="px-3 py-1.5 rounded-lg text-xs font-medium text-slate-600 hover:text-slate-900 cursor-pointer transition-all">
              <i class="fa-regular fa-calendar mr-1"></i>月曆視圖
            </button>
          </div>
        </div>
      </div>

      <!-- Table View Content -->
      <div id="cash-table-wrapper">
        <!-- Status Filter Tabs -->
        <div class="p-3 bg-white border-b border-slate-100 flex flex-wrap items-center gap-1.5">
          <button onclick="setAuditFilter('PENDING_ALL')" id="audit-pill-PENDING_ALL" class="px-3 py-1.5 rounded-lg text-xs font-bold cursor-pointer transition-all">
            待處理 (<span id="count-PENDING_ALL">0</span>)
          </button>
          <button onclick="setAuditFilter('DECLARED_PENDING_BANK')" id="audit-pill-DECLARED_PENDING_BANK" class="px-3 py-1.5 rounded-lg text-xs font-bold cursor-pointer transition-all">
            🟡 已申報待入帳 (<span id="count-DECLARED_PENDING_BANK">0</span>)
          </button>
          <button onclick="setAuditFilter('UNPAID')" id="audit-pill-UNPAID" class="px-3 py-1.5 rounded-lg text-xs font-bold cursor-pointer transition-all">
            🔴 門市未存 (<span id="count-UNPAID">0</span>)
          </button>
          <button onclick="setAuditFilter('SHORT_OVER')" id="audit-pill-SHORT_OVER" class="px-3 py-1.5 rounded-lg text-xs font-bold cursor-pointer transition-all">
            ⚖️ 存款短溢 (<span id="count-SHORT_OVER">0</span>)
          </button>
          <button onclick="setAuditFilter('MATCHED')" id="audit-pill-MATCHED" class="px-3 py-1.5 rounded-lg text-xs font-bold cursor-pointer transition-all">
            🟢 雙審通過 (<span id="count-MATCHED">0</span>)
          </button>
          <button onclick="setAuditFilter('ALL_DAYS')" id="audit-pill-ALL_DAYS" class="px-3 py-1.5 rounded-lg text-xs font-bold cursor-pointer transition-all">
            全部營業日 (<span id="count-ALL_DAYS">0</span>)
          </button>
        </div>

        <div class="overflow-x-auto">
          <table class="w-full text-left text-xs">
            <thead class="bg-slate-50 text-slate-500 font-semibold border-b border-slate-200">
              <tr>
                <th class="py-3 px-4">門市</th>
                <th class="py-3 px-4">營業日期</th>
                <th class="py-3 px-4 text-right">POS 應收現金</th>
                <th class="py-3 px-4">第一層：營運總表 U 欄 (門市申報)</th>
                <th class="py-3 px-4 text-center">短溢差額 (實存-應收)</th>
                <th class="py-3 px-4">第二層：玉山銀行存摺 (銀行入帳)</th>
                <th class="py-3 px-4 text-center">雙層審核結果</th>
                <th class="py-3 px-4 text-center">操作</th>
              </tr>
            </thead>
            <tbody class="divide-y divide-slate-100 text-slate-700" id="dual-audit-tbody">
              <!-- Rendered by JS -->
            </tbody>
          </table>
        </div>
      </div>

      <!-- Calendar View Content -->
      <div id="cash-calendar-wrapper" class="hidden p-4 sm:p-5">
        <div class="flex flex-wrap items-center justify-between gap-3 pb-3 border-b border-slate-100 text-xs">
          <h3 class="font-bold text-slate-800" id="calendar-title-text">營業與存款核對月曆</h3>
          <div class="flex flex-wrap items-center gap-3 text-xs font-medium text-slate-600">
            <span class="flex items-center gap-1.5"><span class="w-2.5 h-2.5 rounded-full bg-emerald-500"></span> 雙審通過</span>
            <span class="flex items-center gap-1.5"><span class="w-2.5 h-2.5 rounded-full bg-amber-400"></span> 已申報待入帳</span>
            <span class="flex items-center gap-1.5"><span class="w-2.5 h-2.5 rounded-full bg-rose-500"></span> 留店未存</span>
            <span class="flex items-center gap-1.5"><span class="w-2.5 h-2.5 rounded-full bg-blue-500"></span> 銀行已收漏填</span>
          </div>
        </div>

        <div id="calendar-day-detail" class="hidden my-3 p-4 rounded-xl bg-slate-800 text-white text-xs shadow-sm"></div>

        <div class="grid grid-cols-7 gap-2 my-2 text-center text-xs font-bold text-slate-400">
          <div class="text-rose-500">日</div>
          <div>一</div>
          <div>二</div>
          <div>三</div>
          <div>四</div>
          <div>五</div>
          <div class="text-indigo-500">六</div>
        </div>
        <div class="grid grid-cols-7 gap-2" id="calendar-grid">
          <!-- Rendered by JS -->
        </div>
      </div>
    </section>
  </div>

  <!-- Tab Pane 2: 外送與金流四步推算 & 平台發票追蹤 -->
  <div id="tab-pane-channels" class="hidden space-y-4">
    <!-- Channel 4-Step Deduction Summary Table -->
    <section class="card-box overflow-hidden">
      <div class="p-4 sm:p-5 border-b border-slate-200 flex flex-col xl:flex-row justify-between items-start xl:items-center gap-3 bg-slate-50/60">
        <div>
          <div class="flex items-center gap-2">
            <span class="w-2.5 h-2.5 rounded-full bg-indigo-600"></span>
            <h2 class="font-bold text-slate-900 text-base">各收款通路「POS 原始額 ➔ 扣費後實收額 ➔ 銀行已存額 ➔ 推算未入帳差額」匯總表</h2>
          </div>
          <p class="text-xs text-slate-500 mt-0.5">直接採用營運總表「實際收款 (扣手續費後淨額)」對比玉山銀行「實際入帳金額」，精準推算各通路待入帳淨額</p>
        </div>
        <!-- Accrual Bridge Mini Summary -->
        <div class="flex flex-wrap items-center gap-3 text-xs bg-white px-3.5 py-2 rounded-xl border border-slate-200">
          <span class="text-slate-500">🏦 存摺當月流水 <b class="text-slate-800 font-mono" id="bridge-passbook">$0</b></span>
          <span class="text-slate-300">|</span>
          <span class="text-slate-500">➖ 代收上月款 <b class="text-amber-600 font-mono" id="bridge-prev">-$0</b></span>
          <span class="text-slate-300">|</span>
          <span class="text-slate-500">📊 屬本月沖銷 <b class="text-emerald-600 font-mono" id="bridge-revenue">$0</b></span>
        </div>
      </div>
      <div class="overflow-x-auto">
        <table class="w-full text-left text-xs">
          <thead class="bg-slate-50 text-slate-500 font-semibold border-b border-slate-200">
            <tr>
              <th class="py-3 px-4">收款通路</th>
              <th class="py-3 px-4 text-right">① POS 原始額</th>
              <th class="py-3 px-4 text-right">平台手續費/抽成</th>
              <th class="py-3 px-4 text-right text-indigo-700">② 扣手續費後實收額 (實收)</th>
              <th class="py-3 px-4 text-right text-emerald-700">③ 玉山銀行已存額 (已存)</th>
              <th class="py-3 px-4 text-right text-rose-600">④ 推算未入帳差額 (② - ③)</th>
              <th class="py-3 px-4">核帳與撥款狀態說明</th>
            </tr>
          </thead>
          <tbody class="divide-y divide-slate-100 text-slate-700" id="platform-cards-container">
            <!-- Rendered by JS -->
          </tbody>
        </table>
      </div>
    </section>

    <!-- Platform Fee Invoice Tracking Table -->
    <section class="card-box overflow-hidden" id="invoice-tracking-section">
      <div class="p-4 sm:p-5 border-b border-slate-200 flex flex-col sm:flex-row justify-between items-start sm:items-center gap-3 bg-slate-50/60">
        <div>
          <div class="flex items-center gap-2">
            <span class="w-2.5 h-2.5 rounded-full bg-blue-600"></span>
            <h2 class="font-bold text-slate-900 text-base">🧾 平台手續費發票（進項扣抵）缺漏追蹤清單</h2>
          </div>
          <p class="text-xs text-slate-500 mt-0.5">自動試算各外送與電子支付平台應取得之手續費發票金額與可扣抵 5% 進項營業稅額，避免漏報多繳稅</p>
        </div>
        <div class="flex flex-wrap items-center gap-2">
          <div class="text-xs font-bold px-3 py-1.5 rounded-xl bg-blue-50 text-blue-700 border border-blue-200" id="invoice-summary-badge">
            預估可扣抵進項稅額合計：$0
          </div>
          <button onclick="openInvoiceUploadModal()" class="px-3.5 py-1.5 rounded-xl bg-indigo-600 hover:bg-indigo-700 text-white font-bold text-xs flex items-center gap-1.5 shadow-sm transition-all cursor-pointer">
            <i class="fa-solid fa-cloud-arrow-up"></i>
            <span>📤 上傳 / 登記進項發票</span>
          </button>
        </div>
      </div>
      <div class="overflow-x-auto">
        <table class="w-full text-left text-xs">
          <thead class="bg-slate-50 text-slate-500 font-semibold border-b border-slate-200">
            <tr>
              <th class="py-3 px-4">門市</th>
              <th class="py-3 px-4">平台名稱 / 廠商統編</th>
              <th class="py-3 px-4 text-center">預估費率</th>
              <th class="py-3 px-4 text-right">本月 POS 營收</th>
              <th class="py-3 px-4 text-right">預估全月手續費(含稅)</th>
              <th class="py-3 px-4 text-right">可扣抵 5% 進項稅額</th>
              <th class="py-3 px-4 text-center">發票沖銷狀態</th>
              <th class="py-3 px-4">已登記進項發票</th>
              <th class="py-3 px-4">會計沖銷與登記操作</th>
            </tr>
          </thead>
          <tbody class="divide-y divide-slate-100 text-slate-700" id="invoice-tracking-tbody">
            <!-- Rendered by JS -->
          </tbody>
        </table>
      </div>
    </section>
  </div>

  <!-- Tab Pane 3: 玉山銀行入帳沖銷流水 -->
  <div id="tab-pane-bank" class="hidden space-y-4">
    <section class="card-box overflow-hidden">
      <div class="p-4 sm:p-5 border-b border-slate-200 flex flex-col md:flex-row justify-between items-start md:items-center gap-3 bg-slate-50/60">
        <div>
          <div class="flex items-center gap-2">
            <span class="w-2.5 h-2.5 rounded-full bg-emerald-600"></span>
            <h2 class="font-bold text-slate-900 text-base">玉山銀行單筆入帳沖銷明細</h2>
          </div>
          <p class="text-xs text-slate-500 mt-0.5">依銀行存摺入帳順序排列，標註每筆款項對應的營業日期與跨月歸屬</p>
        </div>
        <div class="w-full md:w-64 relative">
          <i class="fa-solid fa-magnifying-glass absolute left-3 top-2.5 text-slate-400 text-xs"></i>
          <input type="text" id="search-input" oninput="filterTransactions()" placeholder="搜尋日期、門市、金額..." 
                 class="w-full pl-8 pr-3 py-1.5 bg-white border border-slate-200 rounded-xl text-xs focus:outline-none focus:ring-2 focus:ring-indigo-500">
        </div>
      </div>

      <!-- Category Pills -->
      <div class="px-5 py-2.5 bg-white border-b border-slate-100 flex flex-wrap items-center gap-1.5">
        <button onclick="setCategoryFilter('ALL')" class="cat-pill px-3 py-1 rounded-lg text-xs font-bold bg-indigo-600 text-white cursor-pointer" id="cat-pill-ALL">全部</button>
        <button onclick="setCategoryFilter('現金')" class="cat-pill px-3 py-1 rounded-lg text-xs font-semibold bg-slate-100 text-slate-600 hover:bg-slate-200 cursor-pointer" id="cat-pill-現金">💵 現金存入</button>
        <button onclick="setCategoryFilter('LINE')" class="cat-pill px-3 py-1 rounded-lg text-xs font-semibold bg-slate-100 text-slate-600 hover:bg-slate-200 cursor-pointer" id="cat-pill-LINE">💚 LINE Pay</button>
        <button onclick="setCategoryFilter('Uber')" class="cat-pill px-3 py-1 rounded-lg text-xs font-semibold bg-slate-100 text-slate-600 hover:bg-slate-200 cursor-pointer" id="cat-pill-Uber">🚗 Uber Eats</button>
        <button onclick="setCategoryFilter('Foodpanda')" class="cat-pill px-3 py-1 rounded-lg text-xs font-semibold bg-slate-100 text-slate-600 hover:bg-slate-200 cursor-pointer" id="cat-pill-Foodpanda">🐼 Foodpanda</button>
        <button onclick="setCategoryFilter('偷訂')" class="cat-pill px-3 py-1 rounded-lg text-xs font-semibold bg-slate-100 text-slate-600 hover:bg-slate-200 cursor-pointer" id="cat-pill-偷訂">🛒 偷訂</button>
        <button onclick="setCategoryFilter('CROSS')" class="cat-pill px-3 py-1 rounded-lg text-xs font-semibold bg-slate-100 text-slate-600 hover:bg-slate-200 cursor-pointer" id="cat-pill-CROSS">🔄 跨月款項</button>
      </div>

      <div class="overflow-x-auto">
        <table class="w-full text-left text-xs">
          <thead class="bg-slate-50 text-slate-500 font-semibold border-b border-slate-200">
            <tr>
              <th class="py-3 px-4 w-14 text-center">#</th>
              <th class="py-3 px-4">門市</th>
              <th class="py-3 px-4">銀行入帳日</th>
              <th class="py-3 px-4">通路類別</th>
              <th class="py-3 px-4 text-right">銀行入帳金額</th>
              <th class="py-3 px-4">對應營業日期</th>
              <th class="py-3 px-4">月份歸屬說明</th>
            </tr>
          </thead>
          <tbody class="divide-y divide-slate-100 text-slate-700" id="transaction-tbody">
            <!-- Rendered by JS -->
          </tbody>
        </table>
      </div>
    </section>
  </div>

  <!-- Invoice Upload Modal -->
  <div id="invoice-upload-modal" class="fixed inset-0 z-50 hidden bg-slate-900/60 backdrop-blur-sm flex items-center justify-center p-4">
    <div class="bg-white rounded-2xl shadow-2xl max-w-lg w-full overflow-hidden border border-slate-200">
      <!-- Header -->
      <div class="px-6 py-4 bg-gradient-to-r from-slate-900 via-indigo-950 to-slate-900 text-white flex justify-between items-center">
        <div class="flex items-center gap-2.5">
          <div class="w-8 h-8 rounded-lg bg-indigo-500/20 border border-indigo-400/30 flex items-center justify-center text-indigo-300">
            <i class="fa-solid fa-receipt text-sm"></i>
          </div>
          <div>
            <h3 class="text-sm font-bold">上傳 / 登記平台手續費進項發票</h3>
            <p class="text-[11px] text-slate-300">登記發票號碼與金額，系統將自動扣抵差額並提列 5% 進項營業稅</p>
          </div>
        </div>
        <button onclick="closeInvoiceUploadModal()" class="text-slate-400 hover:text-white text-base cursor-pointer">✕</button>
      </div>

      <!-- Form Body -->
      <form id="invoice-upload-form" onsubmit="handleInvoiceFormSubmit(event)" class="p-6 space-y-4 text-xs">
        <div class="grid grid-cols-2 gap-3">
          <div>
            <label class="block font-bold text-slate-700 mb-1">🏪 門市</label>
            <select id="modal-inv-store" required class="w-full px-3 py-2 bg-slate-50 border border-slate-200 rounded-xl focus:bg-white focus:outline-none focus:ring-2 focus:ring-indigo-500">
              <option value="月宮">月宮 (月下餐飲)</option>
              <option value="隨意小吃">隨意小吃 (月下餐飲)</option>
            </select>
          </div>
          <div>
            <label class="block font-bold text-slate-700 mb-1">🛵 平台 / 廠商</label>
            <select id="modal-inv-platform" required class="w-full px-3 py-2 bg-slate-50 border border-slate-200 rounded-xl focus:bg-white focus:outline-none focus:ring-2 focus:ring-indigo-500">
              <option value="LINE Pay (連加網路)">LINE Pay (連加網路 · 99024139)</option>
              <option value="Uber Eats (優步科技)">Uber Eats (優步科技 · 83118125)</option>
              <option value="Foodpanda (富胖達)">Foodpanda (富胖達 · 53075607)</option>
            </select>
          </div>
        </div>

        <div class="grid grid-cols-2 gap-3">
          <div>
            <label class="block font-bold text-slate-700 mb-1">🏷️ 發票號碼 <span class="text-rose-500">*</span></label>
            <input type="text" id="modal-inv-no" required placeholder="如：AB-12345678" maxlength="12"
                   class="w-full px-3 py-2 bg-slate-50 border border-slate-200 rounded-xl font-mono uppercase focus:bg-white focus:outline-none focus:ring-2 focus:ring-indigo-500">
          </div>
          <div>
            <label class="block font-bold text-slate-700 mb-1">📅 發票開立日期</label>
            <input type="date" id="modal-inv-date" required
                   class="w-full px-3 py-2 bg-slate-50 border border-slate-200 rounded-xl font-mono focus:bg-white focus:outline-none focus:ring-2 focus:ring-indigo-500">
          </div>
        </div>

        <div>
          <label class="block font-bold text-slate-700 mb-1">💰 發票含稅總金額 (新台幣) <span class="text-rose-500">*</span></label>
          <div class="relative">
            <span class="absolute left-3 top-2 text-slate-400 font-bold">$</span>
            <input type="number" id="modal-inv-amount" required min="1" step="1" placeholder="請輸入含稅總金額" oninput="calculateVatSplit()"
                   class="w-full pl-7 pr-3 py-2 bg-slate-50 border border-slate-200 rounded-xl font-mono text-sm font-bold text-slate-900 focus:bg-white focus:outline-none focus:ring-2 focus:ring-indigo-500">
          </div>
          <!-- Real-time VAT computation display -->
          <div class="mt-2 p-2.5 rounded-xl bg-indigo-50 border border-indigo-100 flex items-center justify-between text-[11px] text-indigo-900">
            <span>未稅銷售額：<b id="modal-inv-untaxed" class="font-mono">$0</b></span>
            <span class="text-indigo-300">|</span>
            <span>可扣抵 5% 進項稅額：<b id="modal-inv-tax" class="font-mono text-blue-600">$0</b></span>
          </div>
        </div>

        <div>
          <label class="block font-bold text-slate-700 mb-1">📎 發票或對帳單檔案 (選填，支援 PDF / Excel / 圖檔)</label>
          <div id="drop-zone" onclick="document.getElementById('modal-inv-file').click()" 
               class="border-2 border-dashed border-slate-200 hover:border-indigo-400 bg-slate-50/50 hover:bg-indigo-50/30 rounded-xl p-4 text-center cursor-pointer transition-all">
            <input type="file" id="modal-inv-file" class="hidden" accept=".pdf,.xlsx,.xls,.csv,.png,.jpg,.jpeg" onchange="handleFileSelected(this)">
            <div id="drop-zone-text" class="text-slate-500 space-y-1">
              <i class="fa-solid fa-cloud-arrow-up text-xl text-indigo-500"></i>
              <p class="font-medium text-[11px]">點擊或拖曳檔案至此處上傳</p>
              <p class="text-[10px] text-slate-400">支援格式：PDF、Excel (XLSX/CSV)、電子發票截圖 (PNG/JPG)</p>
            </div>
            <div id="drop-zone-selected" class="hidden text-slate-700 items-center justify-between bg-white px-3 py-2 rounded-lg border border-slate-200">
              <span id="selected-file-name" class="font-mono truncate text-[11px]">filename.pdf</span>
              <button type="button" onclick="event.stopPropagation(); clearSelectedFile()" class="text-rose-500 hover:text-rose-700 text-xs ml-2 cursor-pointer">✕ 移除</button>
            </div>
          </div>
        </div>

        <div>
          <label class="block font-bold text-slate-700 mb-1">📝 備註 (選填)</label>
          <input type="text" id="modal-inv-note" placeholder="例如：10月手續費發票 / 雙十活動扣款"
                 class="w-full px-3 py-2 bg-slate-50 border border-slate-200 rounded-xl focus:bg-white focus:outline-none focus:ring-2 focus:ring-indigo-500">
        </div>

        <div id="modal-error-msg" class="hidden p-2.5 rounded-xl bg-rose-50 border border-rose-200 text-rose-700 text-[11px]"></div>

        <div class="pt-3 border-t border-slate-100 flex items-center justify-end gap-2.5">
          <button type="button" onclick="closeInvoiceUploadModal()" class="px-4 py-2 rounded-xl border border-slate-200 text-slate-600 hover:bg-slate-100 font-bold transition-all cursor-pointer">
            取消
          </button>
          <button type="submit" id="modal-submit-btn" class="px-5 py-2 rounded-xl bg-indigo-600 hover:bg-indigo-700 text-white font-bold flex items-center gap-2 shadow-sm transition-all cursor-pointer">
            <i class="fa-solid fa-check"></i>
            <span>💾 確認登記並即時核銷</span>
          </button>
        </div>
      </form>
    </div>
  </div>

  <!-- Script -->
  <script>
    let allMonthsData = {json_payload_str};
    let currentMonth = '{target_month}';
    let currentStore = 'ALL';
    let currentCategoryFilter = 'ALL';
    let currentAuditFilter = 'PENDING_ALL';
    let currentDailyCashMap = {{}};
    let currentMainTab = 'cash';
    let currentCashView = 'table';

    function switchMainTab(tabKey) {{
      currentMainTab = tabKey;
      ['cash', 'channels', 'bank'].forEach(k => {{
        const pane = document.getElementById('tab-pane-' + k);
        const tabBtn = document.getElementById('main-tab-btn-' + k);
        if (pane) {{
          if (k === tabKey) pane.classList.remove('hidden');
          else pane.classList.add('hidden');
        }}
        if (tabBtn) {{
          if (k === tabKey) {{
            tabBtn.className = 'main-nav-tab px-4 py-2 rounded-xl text-xs sm:text-sm font-bold bg-indigo-600 text-white shadow-sm flex items-center gap-2 cursor-pointer transition-all';
          }} else {{
            tabBtn.className = 'main-nav-tab px-4 py-2 rounded-xl text-xs sm:text-sm font-semibold bg-slate-100 text-slate-600 hover:bg-slate-200 flex items-center gap-2 cursor-pointer transition-all';
          }}
        }}
      }});
    }}

    function toggleCashView(viewMode) {{
      currentCashView = viewMode;
      const tblSec = document.getElementById('cash-table-wrapper');
      const calSec = document.getElementById('cash-calendar-wrapper');
      const btnTbl = document.getElementById('btn-cash-view-table');
      const btnCal = document.getElementById('btn-cash-view-calendar');
      if (viewMode === 'table') {{
        if (tblSec) tblSec.classList.remove('hidden');
        if (calSec) calSec.classList.add('hidden');
        if (btnTbl) btnTbl.className = 'px-3 py-1.5 rounded-lg text-xs font-bold bg-white text-indigo-700 shadow-xs cursor-pointer transition-all';
        if (btnCal) btnCal.className = 'px-3 py-1.5 rounded-lg text-xs font-medium text-slate-600 hover:text-slate-900 cursor-pointer transition-all';
      }} else {{
        if (tblSec) tblSec.classList.add('hidden');
        if (calSec) calSec.classList.remove('hidden');
        if (btnCal) btnCal.className = 'px-3 py-1.5 rounded-lg text-xs font-bold bg-white text-indigo-700 shadow-xs cursor-pointer transition-all';
        if (btnTbl) btnTbl.className = 'px-3 py-1.5 rounded-lg text-xs font-medium text-slate-600 hover:text-slate-900 cursor-pointer transition-all';
      }}
    }}

    function updateHeaderTitle() {{
      const storeLabel = currentStore === 'ALL' ? '全部門市' : currentStore;
      const fullTitle = `餐飲集團-${{storeLabel}}`;
      const headerEl = document.getElementById('header-main-title');
      if (headerEl) headerEl.textContent = fullTitle;
      const footerEl = document.getElementById('footer-title');
      if (footerEl) footerEl.textContent = `${{fullTitle}} · 自動化財務對帳系統`;
      if (typeof document !== 'undefined' && 'title' in document) {{
        document.title = `${{fullTitle}} · 財務對帳看板`;
      }}
    }}

    function showToast(msg) {{
      const t = document.getElementById('sync-toast');
      const txt = document.getElementById('sync-toast-text');
      if (!t || !txt) return;
      txt.textContent = msg;
      t.classList.remove('hidden');
      t.classList.add('flex');
      setTimeout(() => {{
        t.classList.add('hidden');
        t.classList.remove('flex');
      }}, 2200);
    }}

    async function refreshLiveData(manual = false) {{
      const icon = document.getElementById('refresh-icon');
      const textEl = document.getElementById('refresh-btn-text');
      if (icon) icon.classList.add('fa-spin');
      if (textEl && manual) textEl.textContent = '抓取中...';

      let liveSynced = false;

      // 1. 優先嘗試連線本機 Python 服務 (localhost:8080/api/sync) 觸發重新抓取 Google 試算表與即時計算
      try {{
        const controller = new AbortController();
        const timeoutId = setTimeout(() => controller.abort(), 35000);
        const apiHost = (typeof window !== 'undefined' && window.location && window.location.hostname === 'localhost')
          ? '' : 'http://localhost:8080';

        const apiRes = await fetch(apiHost + '/api/sync', {{
          method: 'POST',
          headers: {{ 'Content-Type': 'application/json' }},
          body: JSON.stringify({{ month: currentMonth }}),
          signal: controller.signal
        }});
        clearTimeout(timeoutId);

        if (apiRes.ok) {{
          const resJson = await apiRes.json();
          if (resJson && resJson.success && resJson.data) {{
            allMonthsData = resJson.data;
            const uTime = resJson.updated_at || '';
            const el1 = document.getElementById('header-updated-at');
            if (el1) el1.textContent = uTime;
            const el2 = document.getElementById('top-updated-at');
            if (el2) el2.textContent = uTime;
            renderDashboard();
            liveSynced = true;
            if (manual) showToast('⚡ 成功重新抓取 Google 試算表，時間已更新為當下！');
          }}
        }}
      }} catch (err) {{
        // 本機服務未啟動時，自動降級為抓取靜態雲端快照
      }}

      // 2. 若本機 API 未回應（例如外出以手機或 GitHub Pages 瀏覽），則讀取雲端/本地已產出的 dashboard_data.json
      if (!liveSynced) {{
        try {{
          const res = await fetch('./dashboard_data.json?t=' + Date.now(), {{ cache: 'no-store' }});
          if (res.ok) {{
            const freshData = await res.json();
            if (freshData && Object.keys(freshData).length > 0) {{
              allMonthsData = freshData;
              if (allMonthsData[currentMonth]) {{
                const uTime = allMonthsData[currentMonth].updated_at || '';
                const el1 = document.getElementById('header-updated-at');
                if (el1) el1.textContent = uTime;
                const el2 = document.getElementById('top-updated-at');
                if (el2) el2.textContent = uTime;
              }}
              renderDashboard();
              if (manual) showToast('☁️ 已同步雲端最新結算數據');
            }}
          }}
        }} catch (e) {{
          if (manual) {{
            renderDashboard();
            showToast('畫面資料已重新整理');
          }}
        }}
      }}

      if (icon) setTimeout(() => icon.classList.remove('fa-spin'), 300);
      if (textEl) setTimeout(() => textEl.textContent = '即時更新', 300);
    }}

    function jumpAndFilterAudit(statusFilter) {{
      switchMainTab('cash');
      toggleCashView('table');
      setAuditFilter(statusFilter);
      const sec = document.getElementById('alert-section');
      if (sec) sec.scrollIntoView({{ behavior: 'smooth', block: 'start' }});
    }}

    function scrollToInvoiceSection() {{
      switchMainTab('channels');
      const sec = document.getElementById('invoice-tracking-section');
      if (sec) sec.scrollIntoView({{ behavior: 'smooth', block: 'start' }});
    }}

    function jumpToBankCategory(cat) {{
      switchMainTab('bank');
      setCategoryFilter(cat);
      const sec = document.getElementById('tab-pane-bank');
      if (sec) sec.scrollIntoView({{ behavior: 'smooth', block: 'start' }});
    }}

    function toggleDownloadMenu() {{
      const m = document.getElementById('download-menu');
      if (m) m.classList.toggle('hidden');
    }}

    function updateDownloadLinks() {{
      const mClean = currentMonth.replace('-', '');
      const prefix = (typeof window !== 'undefined' && window.location && window.location.pathname.includes('/output/')) ? '.' : './output';
      const l1 = document.getElementById('dl-summary-excel');
      if (l1) l1.href = `${{prefix}}/本月對帳差額與應收餘額表_${{mClean}}.xlsx`;
      const l2 = document.getElementById('dl-bank-voucher');
      if (l2) l2.href = `${{prefix}}/文中傳票維護匯入_全部合併_玉山銀行沖銷_${{mClean}}.xlsx`;
      const l3 = document.getElementById('dl-pos-voucher');
      if (l3) l3.href = `${{prefix}}/文中傳票維護匯入_全部合併_${{mClean}}.xlsx`;
    }}

    function closeLineModal() {{
      const modal = document.getElementById('line-modal');
      if (modal) {{
        modal.classList.add('hidden');
        modal.classList.remove('flex');
      }}
    }}

    function copyTextareaAgain() {{
      const ta = document.getElementById('line-summary-textarea');
      if (!ta) return;
      ta.select();
      if (typeof navigator !== 'undefined' && navigator.clipboard) {{
        navigator.clipboard.writeText(ta.value);
      }}
      showToast('已複製到剪貼簿！可直接貼上 LINE');
    }}

    function copyLineSummary() {{
      const rawData = allMonthsData[currentMonth];
      if (!rawData) return;

      const storeLabel = currentStore === 'ALL' ? '全部門市' : currentStore;
      const storesToInclude = currentStore === 'ALL' ? Object.keys(rawData.stores || {{}}) : [currentStore];

      let unpaidList = [];
      let declaredList = [];
      let shortOverList = [];
      let seenShortOverPeriods = {{}};
      let unpaidTotal = 0;
      let declaredTotal = 0;

      const chTotals = {{
        '1130-01': {{ label: '💵 現金', pos: 0, netRec: 0, bankDep: 0, netUnpaid: 0 }},
        '1130-02': {{ label: '💚 LINE Pay', pos: 0, netRec: 0, bankDep: 0, netUnpaid: 0 }},
        '1130-03': {{ label: '🚗 Uber Eats', pos: 0, netRec: 0, bankDep: 0, netUnpaid: 0 }},
        '1130-04': {{ label: '🐼 Foodpanda', pos: 0, netRec: 0, bankDep: 0, netUnpaid: 0 }}
      }};

      storesToInclude.forEach(sk => {{
        const s = (rawData.stores || {{}})[sk];
        if (!s) return;
        const sName = (s.name || sk).replace(' (月下餐飲)', '');

        Object.entries(s.channels || {{}}).forEach(([code, c]) => {{
          if (chTotals[code]) {{
            const posVal = c.pos_sales || 0;
            const netRecVal = (c.net_after_fee !== undefined) ? c.net_after_fee : posVal;
            const bankDepVal = (c.bank_deposited !== undefined) ? c.bank_deposited : (c.bank_cleared || 0);
            const unpaidVal = (c.net_unpaid_diff !== undefined) ? c.net_unpaid_diff : Math.max(0, netRecVal - bankDepVal);
            chTotals[code].pos += posVal;
            chTotals[code].netRec += netRecVal;
            chTotals[code].bankDep += bankDepVal;
            chTotals[code].netUnpaid += unpaidVal;
          }}
        }});

        (s.daily_cash || []).forEach(d => {{
          if ((d.cash || 0) <= 0) return;
          if (d.audit_status === 'UNPAID') {{
            unpaidTotal += d.cash;
            unpaidList.push(`• ${{sName}}｜${{d.date.slice(5)}} 營業現金 $${{Math.round(d.cash).toLocaleString()}}（總表U欄未填）`);
          }} else if (d.audit_status === 'DECLARED_PENDING_BANK') {{
            declaredTotal += d.cash;
            const note = (d.declared_note || '').replace('申報於 ', '').replace(/2026-/g, '').replace('2026/', '');
            declaredList.push(`• ${{sName}}｜${{d.date.slice(5)}} 營業現金 $${{Math.round(d.cash).toLocaleString()}}（${{note}}）`);
          }}
          const so = d.short_over || 0;
          if (Math.abs(so) >= 1) {{
            const pKey = `${{sName}}_${{d.period || d.date}}`;
            if (!seenShortOverPeriods[pKey]) {{
              seenShortOverPeriods[pKey] = true;
              const pDisp = (d.period || d.date).replace(/2026-/g, '');
              const gSum = Math.round(d.group_pos_sum || d.cash).toLocaleString();
              const dAmt = Math.round(d.declared_amt || 0).toLocaleString();
              const soTag = so < 0 ? `⚠️ 短少 -$${{Math.abs(Math.round(so)).toLocaleString()}}` : `➕ 溢報 +$${{Math.round(so).toLocaleString()}}`;
              shortOverList.push(`• ${{sName}}｜${{pDisp}} 應收合計 $${{gSum}}，申報存入 $${{dAmt}}（${{soTag}}）`);
            }}
          }}
        }});
      }});

      let lines = [
        `📊【餐飲集團-${{storeLabel}}】${{currentMonth}} 實收 vs 已存 推算與催帳摘要`,
        `🕒 更新時間：${{rawData.updated_at || ''}}`,
        `━━━━━━━━━━━━━━`,
        `📌【各通路 實收(扣手續費後) vs 已存(銀行入帳) 推算】`
      ];

      Object.values(chTotals).forEach(ch => {{
        if (ch.pos > 0 || ch.bankDep > 0) {{
          lines.push(`${{ch.label}}：實收 $${{Math.round(ch.netRec).toLocaleString()}} / 已存 $${{Math.round(ch.bankDep).toLocaleString()}} ➔ 未入帳差額 $${{Math.round(ch.netUnpaid).toLocaleString()}}`);
        }}
      }});

      lines.push(`━━━━━━━━━━━━━━`);

      if (unpaidList.length > 0) {{
        lines.push(`🔴 門市現金未存（請盡速存入並填寫總表U欄）：共 $${{Math.round(unpaidTotal).toLocaleString()}}`);
        lines.push(...unpaidList);
      }} else {{
        lines.push(`🔴 門市現金未存：$0（全數已申報或已入帳）`);
      }}

      lines.push('');

      if (declaredList.length > 0) {{
        lines.push(`🟡 門市現金已申報・待玉山銀行入帳：共 $${{Math.round(declaredTotal).toLocaleString()}}`);
        lines.push(...declaredList);
      }} else {{
        lines.push(`🟡 門市現金已申報待入帳：$0`);
      }}

      if (shortOverList.length > 0) {{
        lines.push('');
        lines.push(`⚖️ 門市存款短溢差額提醒（請核對實存金額）：`);
        lines.push(...shortOverList);
      }}

      lines.push(`━━━━━━━━━━━━━━`);
      lines.push(`🔗 線上即時看板：https://fat870920pon-coder.github.io/-/`);

      const summaryText = lines.join('\\n');
      const ta = document.getElementById('line-summary-textarea');
      if (ta) ta.value = summaryText;

      if (typeof navigator !== 'undefined' && navigator.clipboard) {{
        navigator.clipboard.writeText(summaryText).catch(() => {{}});
      }}

      const modal = document.getElementById('line-modal');
      if (modal) {{
        modal.classList.remove('hidden');
        modal.classList.add('flex');
      }}
      showToast('已自動複製 LINE 催帳摘要！');
    }}

    function updateAuditPillStyles() {{
      const activeMap = {{
        'PENDING_ALL': 'bg-slate-900 text-white shadow-xs',
        'DECLARED_PENDING_BANK': 'bg-amber-500 text-white shadow-xs',
        'UNPAID': 'bg-rose-600 text-white shadow-xs',
        'SHORT_OVER': 'bg-purple-600 text-white shadow-xs',
        'MATCHED': 'bg-emerald-600 text-white shadow-xs',
        'ALL_DAYS': 'bg-indigo-600 text-white shadow-xs'
      }};
      const inactiveClass = 'bg-white text-slate-600 hover:bg-slate-100 border border-slate-200';
      Object.keys(activeMap).forEach(k => {{
        const btn = document.getElementById('audit-pill-' + k);
        if (btn) {{
          btn.className = 'px-3 py-1.5 rounded-lg text-xs font-bold cursor-pointer transition-all ' + (k === currentAuditFilter ? activeMap[k] : inactiveClass);
        }}
      }});
    }}

    function setAuditFilter(filterKey) {{
      currentAuditFilter = filterKey;
      updateAuditPillStyles();
      const rawData = allMonthsData[currentMonth];
      if (rawData) renderAlertsAndAuditTable(rawData);
    }}

    function setCategoryFilter(cat) {{
      currentCategoryFilter = cat;
      document.querySelectorAll('.cat-pill').forEach(btn => {{
        btn.className = 'cat-pill px-3 py-1 rounded-lg text-xs font-semibold bg-slate-100 text-slate-600 hover:bg-slate-200 cursor-pointer';
      }});
      const activeBtn = document.getElementById('cat-pill-' + cat);
      if (activeBtn) {{
        activeBtn.className = 'cat-pill px-3 py-1 rounded-lg text-xs font-bold bg-indigo-600 text-white cursor-pointer';
      }}
      filterTransactions();
    }}

    function switchMonth(selectedMonth) {{
      if (allMonthsData[selectedMonth]) {{
        currentMonth = selectedMonth;
        const uTime = allMonthsData[currentMonth].updated_at || '';
        const el1 = document.getElementById('header-updated-at');
        if (el1) el1.textContent = uTime;
        const el2 = document.getElementById('top-updated-at');
        if (el2) el2.textContent = uTime;
        updateDownloadLinks();
        renderDashboard();
      }}
    }}

    function switchStore(storeKey) {{
      currentStore = storeKey;
      document.querySelectorAll('.store-tab').forEach(tab => {{
        tab.className = 'store-tab px-3.5 py-1.5 rounded-lg text-xs font-bold transition-all text-slate-600 hover:text-slate-900 cursor-pointer';
      }});
      const activeTab = document.getElementById('tab-' + storeKey);
      if (activeTab) {{
        activeTab.className = 'store-tab px-3.5 py-1.5 rounded-lg text-xs font-bold transition-all bg-white text-indigo-600 shadow-xs cursor-pointer';
      }}
      updateHeaderTitle();
      renderDashboard();
    }}

    function renderDashboard() {{
      const rawData = allMonthsData[currentMonth];
      if (!rawData) return;

      const uTime = rawData.updated_at || '';
      const el1 = document.getElementById('header-updated-at');
      if (el1) el1.textContent = uTime;
      const el2 = document.getElementById('top-updated-at');
      if (el2) el2.textContent = uTime;

      updateHeaderTitle();
      updateDownloadLinks();
      updateAuditPillStyles();
      try {{ renderBridge(rawData); }} catch (e) {{ console.error(e); }}
      try {{ renderKPI(rawData); }} catch (e) {{ console.error(e); }}
      try {{ renderAlertsAndAuditTable(rawData); }} catch (e) {{ console.error(e); }}
      try {{ renderPlatforms(rawData); }} catch (e) {{ console.error(e); }}
      try {{ renderInvoiceTracking(rawData); }} catch (e) {{ console.error(e); }}
      try {{ renderCalendar(rawData); }} catch (e) {{ console.error(e); }}
      try {{ filterTransactions(rawData); }} catch (e) {{ console.error(e); }}
    }}

    function renderBridge(rawData) {{
      const b = rawData.bridge || {{}};
      const passbookVal = b.bank_passbook_total || 0;
      const prevVal = b.collected_for_prev_month || 0;
      const nextVal = b.collected_in_next_month || 0;
      const revVal = b.accrual_cleared_for_this_month || (passbookVal - prevVal + nextVal);

      const elPassbook = document.getElementById('bridge-passbook');
      if (elPassbook) elPassbook.textContent = '$' + Math.round(passbookVal).toLocaleString();

      const elPrev = document.getElementById('bridge-prev');
      if (elPrev) elPrev.textContent = '-$' + Math.round(prevVal).toLocaleString();

      const elRev = document.getElementById('bridge-revenue');
      if (elRev) elRev.textContent = '$' + Math.round(revVal).toLocaleString();

      const elNote = document.getElementById('kpi-bridge-note');
      if (elNote) {{
        elNote.textContent = prevVal > 0
          ? `存摺實收 $${{Math.round(passbookVal).toLocaleString()}} (已扣上月跨月款 $${{Math.round(prevVal).toLocaleString()}})`
          : `存摺當月實收全數歸屬本月`;
      }}
    }}

    function renderKPI(rawData) {{
      let pos = 0, netRec = 0, bankDep = 0, netUnpaid = 0, estFee = 0, estVat = 0;
      const storesToInclude = currentStore === 'ALL' ? Object.keys(rawData.stores || {{}}) : [currentStore];
      storesToInclude.forEach(sk => {{
        const s = (rawData.stores || {{}})[sk];
        if (!s) return;
        Object.entries(s.channels || {{}}).forEach(([code, c]) => {{
          const cPos = c.pos_sales || 0;
          const cNetRec = (c.net_after_fee !== undefined) ? c.net_after_fee : cPos;
          const cBankDep = (c.bank_deposited !== undefined) ? c.bank_deposited : (c.bank_cleared || 0);
          const cUnpaid = (c.net_unpaid_diff !== undefined) ? c.net_unpaid_diff : Math.max(0, cNetRec - cBankDep);
          pos += cPos;
          netRec += cNetRec;
          bankDep += cBankDep;
          netUnpaid += cUnpaid;
          estFee += Math.max(0, cPos - cNetRec);
        }});
        (s.invoice_tracking || []).forEach(it => {{
          estVat += (it.est_vat_5pct || 0);
        }});
      }});

      const elPos = document.getElementById('kpi-pos');
      if (elPos) elPos.textContent = '$ ' + Math.round(pos).toLocaleString();

      const elNetRec = document.getElementById('kpi-net-rec');
      if (elNetRec) elNetRec.textContent = '$ ' + Math.round(netRec).toLocaleString();

      const elFeeNote = document.getElementById('kpi-fee-note');
      if (elFeeNote) elFeeNote.textContent = `已扣除各平台手續費合計 -$${{Math.round(estFee).toLocaleString()}}`;

      const elBank = document.getElementById('kpi-bank');
      if (elBank) elBank.textContent = '$ ' + Math.round(bankDep).toLocaleString();

      const elInv = document.getElementById('kpi-inv');
      if (elInv) elInv.textContent = '$ ' + Math.round(estFee).toLocaleString();

      const elPending = document.getElementById('kpi-pending-cash');
      if (elPending) elPending.textContent = '$ ' + Math.round(netUnpaid).toLocaleString();

      const elVat = document.getElementById('kpi-vat-sub');
      if (elVat) elVat.textContent = `🧾 可扣抵 5% 稅額 $${{Math.round(estVat).toLocaleString()}} (點此追蹤發票)`;
    }}

    function renderAlertsAndAuditTable(rawData) {{
      const tbody = document.getElementById('dual-audit-tbody');
      if (!tbody) return;

      const storesToInclude = currentStore === 'ALL' ? Object.keys(rawData.stores || {{}}) : [currentStore];
      let allDailyRows = [];
      storesToInclude.forEach(sk => {{
        const s = (rawData.stores || {{}})[sk];
        if (!s) return;
        (s.daily_cash || []).forEach(d => {{
          if ((d.cash || 0) > 0) {{
            allDailyRows.push({{
              storeKey: sk,
              storeName: (s.name || sk).replace(' (月下餐飲)', ''),
              ...d
            }});
          }}
        }});
      }});

      allDailyRows.sort((a, b) => a.date.localeCompare(b.date));

      const countDeclared = allDailyRows.filter(r => r.audit_status === 'DECLARED_PENDING_BANK').length;
      const countUnpaid = allDailyRows.filter(r => r.audit_status === 'UNPAID').length;
      const countShortOver = allDailyRows.filter(r => Math.abs(r.short_over || 0) >= 1).length;
      const countMatched = allDailyRows.filter(r => r.audit_status === 'MATCHED' || r.audit_status === 'BANK_ONLY').length;
      const countPendingAll = countDeclared + countUnpaid;

      const setCnt = (id, val) => {{ const el = document.getElementById(id); if (el) el.textContent = val; }};
      setCnt('count-PENDING_ALL', countPendingAll);
      setCnt('count-DECLARED_PENDING_BANK', countDeclared);
      setCnt('count-UNPAID', countUnpaid);
      setCnt('count-SHORT_OVER', countShortOver);
      setCnt('count-MATCHED', countMatched);
      setCnt('count-ALL_DAYS', allDailyRows.length);

      document.getElementById('alert-title-text').textContent =
        `門市現金雙層審核明細（待處理 ${{countPendingAll}} 筆）`;

      const tableRows = allDailyRows.filter(r => {{
        if (currentAuditFilter === 'PENDING_ALL') return r.audit_status === 'DECLARED_PENDING_BANK' || r.audit_status === 'UNPAID';
        if (currentAuditFilter === 'DECLARED_PENDING_BANK') return r.audit_status === 'DECLARED_PENDING_BANK';
        if (currentAuditFilter === 'UNPAID') return r.audit_status === 'UNPAID';
        if (currentAuditFilter === 'SHORT_OVER') return Math.abs(r.short_over || 0) >= 1;
        if (currentAuditFilter === 'MATCHED') return r.audit_status === 'MATCHED' || r.audit_status === 'BANK_ONLY';
        return true;
      }});

      if (tableRows.length === 0) {{
        tbody.innerHTML = `<tr><td colspan="8" class="py-8 text-center text-slate-400">此篩選條件下無紀錄，可點選右上方「全部營業日」查看完整對照</td></tr>`;
        return;
      }}

      const trsHtml = tableRows.map(r => {{
        const st = r.audit_status;
        const so = r.short_over || 0;
        const rowBg = so <= -1 ? 'bg-rose-50/60 hover:bg-rose-50/90' :
                      st === 'UNPAID' ? 'bg-rose-50/40 hover:bg-rose-50/80' :
                      st === 'DECLARED_PENDING_BANK' ? 'bg-amber-50/40 hover:bg-amber-50/80' :
                      'hover:bg-slate-50';

        // 第一層文字精簡化
        let tier1Html = '';
        if (r.has_declared) {{
          const cleanNote = (r.declared_note || '').replace('申報於 ', '').replace('2026/', '').replace(/2026-/g, '');
          tier1Html = `<div class="flex items-center gap-1.5">
            <span class="px-1.5 py-0.5 rounded bg-emerald-100 text-emerald-800 font-bold text-[11px]">已申報</span>
            <span class="text-slate-700 font-medium">${{cleanNote}}</span>
          </div>`;
        }} else {{
          tier1Html = `<div class="flex items-center gap-1.5">
            <span class="px-1.5 py-0.5 rounded bg-rose-100 text-rose-700 font-bold text-[11px]">未申報</span>
            <span class="text-slate-400">總表 U 欄空白</span>
          </div>`;
        }}

        // 短溢差額徽章
        let shortOverHtml = '';
        if (!r.has_declared && !r.bank_cleared) {{
          shortOverHtml = '<span class="text-slate-300">-</span>';
        }} else if (so <= -1) {{
          const gSum = Math.round(r.group_pos_sum || r.cash).toLocaleString();
          shortOverHtml = `<div>
            <span class="inline-block px-2 py-0.5 rounded bg-rose-100 text-rose-800 border border-rose-300 font-bold font-mono text-[11px]">⚠️ 短少 -$${{Math.abs(Math.round(so)).toLocaleString()}}</span>
            <div class="text-[10px] text-rose-600 mt-0.5">應收合 $${{gSum}}</div>
          </div>`;
        }} else if (so >= 1) {{
          const gSum = Math.round(r.group_pos_sum || r.cash).toLocaleString();
          shortOverHtml = `<div>
            <span class="inline-block px-2 py-0.5 rounded bg-indigo-50 text-indigo-700 border border-indigo-200 font-bold font-mono text-[11px]">➕ 溢填 +$${{Math.round(so).toLocaleString()}}</span>
            <div class="text-[10px] text-slate-400 mt-0.5">應收合 $${{gSum}}</div>
          </div>`;
        }} else {{
          shortOverHtml = '<span class="text-emerald-600 font-semibold text-[11px]">✓ 相符</span>';
        }}

        // 第二層文字精簡化
        let tier2Html = '';
        if (r.bank_cleared) {{
          const bDateShort = (r.bank_date || r.deposit_date || '').replace('2026-', '').replace('2026/', '');
          tier2Html = `<div class="flex items-center gap-1.5">
            <span class="px-1.5 py-0.5 rounded bg-emerald-100 text-emerald-800 font-bold text-[11px]">已入帳</span>
            <span class="text-slate-700 font-mono">${{bDateShort}} 實收 $${{Math.round(r.bank_amt || r.cash).toLocaleString()}}</span>
          </div>`;
        }} else {{
          tier2Html = `<div class="flex items-center gap-1.5">
            <span class="px-1.5 py-0.5 rounded bg-slate-100 text-slate-600 font-semibold text-[11px]">待入帳</span>
            <span class="text-slate-400">存摺尚無入帳紀錄</span>
          </div>`;
        }}

        let statusBadge = '';
        if (st === 'MATCHED') {{
          statusBadge = '<span class="inline-block px-2.5 py-1 rounded-full bg-emerald-100 text-emerald-800 font-bold text-[11px]">🟢 雙審通過</span>';
        }} else if (st === 'DECLARED_PENDING_BANK') {{
          statusBadge = '<span class="inline-block px-2.5 py-1 rounded-full bg-amber-100 text-amber-900 font-bold text-[11px]">🟡 已申報・待入帳</span>';
        }} else if (st === 'BANK_ONLY') {{
          statusBadge = '<span class="inline-block px-2.5 py-1 rounded-full bg-blue-100 text-blue-800 font-bold text-[11px]">🔵 銀行已收・漏填</span>';
        }} else {{
          statusBadge = '<span class="inline-block px-2.5 py-1 rounded-full bg-rose-100 text-rose-800 font-bold text-[11px]">🔴 留店未存</span>';
        }}

        return `
          <tr class="${{rowBg}} transition-colors">
            <td class="py-3 px-4 font-bold text-slate-900 whitespace-nowrap">${{r.storeName}}</td>
            <td class="py-3 px-4 font-mono text-slate-600 whitespace-nowrap">${{r.date.slice(5)}}</td>
            <td class="py-3 px-4 text-right font-mono font-bold text-slate-900 whitespace-nowrap">$ ${{Math.round(r.cash).toLocaleString()}}</td>
            <td class="py-3 px-4">${{tier1Html}}</td>
            <td class="py-3 px-4 text-center whitespace-nowrap">${{shortOverHtml}}</td>
            <td class="py-3 px-4">${{tier2Html}}</td>
            <td class="py-3 px-4 text-center whitespace-nowrap">${{statusBadge}}</td>
            <td class="py-3 px-4 text-center whitespace-nowrap">
              <button onclick="scrollToDate('${{r.date}}')" class="text-indigo-600 hover:text-indigo-800 font-semibold cursor-pointer">查看</button>
            </td>
          </tr>
        `;
      }});
      tbody.innerHTML = trsHtml.join('');
    }}

    function renderPlatforms(rawData) {{
      const topCardsEl = document.getElementById('top-channel-cards');
      const tableBodyEl = document.getElementById('platform-cards-container');

      const channelSummary = {{
        '1130-01': {{
          name: '門市營業現金',
          shortName: '現金',
          cardTitle: '💵 實收現金 / 已存現金',
          recLabel: '實收現金 (門市營業)',
          depLabel: '已存現金 (玉山入帳)',
          diffLabel: '推算未入帳現金',
          icon: 'fa-solid fa-money-bill-wave text-amber-500',
          borderAccent: 'border-t-4 border-t-amber-500',
          bgSoft: 'bg-amber-50/50',
          catKey: '現金',
          pos: 0, netRec: 0, bankDep: 0, netUnpaid: 0, fee: 0, feeRateStr: '免手續費',
          unpaidCash: 0, declaredCash: 0
        }},
        '1130-02': {{
          name: 'LINE Pay',
          shortName: 'LINE Pay',
          cardTitle: '💚 實收 LINE Pay / 已存 LINE Pay',
          recLabel: '實收 LINE Pay (扣費後)',
          depLabel: '已存 LINE Pay (玉山入帳)',
          diffLabel: '推算未入帳 LINE Pay',
          icon: 'fa-brands fa-line text-emerald-500',
          borderAccent: 'border-t-4 border-t-emerald-500',
          bgSoft: 'bg-emerald-50/40',
          catKey: 'LINE',
          pos: 0, netRec: 0, bankDep: 0, netUnpaid: 0, fee: 0, feeRateStr: '約 2.31% (總表實收欄)'
        }},
        '1130-03': {{
          name: 'Uber Eats',
          shortName: 'Uber Eats',
          cardTitle: '🚗 實收 Uber Eats / 已存 Uber Eats',
          recLabel: '實收 Uber Eats (扣費後)',
          depLabel: '已存 Uber Eats (玉山入帳)',
          diffLabel: '推算未入帳 Uber Eats',
          icon: 'fa-solid fa-car-side text-slate-700',
          borderAccent: 'border-t-4 border-t-slate-700',
          bgSoft: 'bg-slate-50/80',
          catKey: 'Uber',
          pos: 0, netRec: 0, bankDep: 0, netUnpaid: 0, fee: 0, feeRateStr: '約 33% 平台費'
        }},
        '1130-04': {{
          name: 'Foodpanda',
          shortName: 'Foodpanda',
          cardTitle: '🐼 實收 Foodpanda / 已存 Foodpanda',
          recLabel: '實收 Foodpanda (扣費後)',
          depLabel: '已存 Foodpanda (玉山入帳)',
          diffLabel: '推算未入帳 Foodpanda',
          icon: 'fa-solid fa-box text-pink-500',
          borderAccent: 'border-t-4 border-t-pink-500',
          bgSoft: 'bg-pink-50/40',
          catKey: 'Foodpanda',
          pos: 0, netRec: 0, bankDep: 0, netUnpaid: 0, fee: 0, feeRateStr: '約 33.6% 平台費'
        }},
        '1130-05': {{
          name: '偷訂科技',
          shortName: '偷訂科技',
          cardTitle: '🛒 實收偷訂 / 已存偷訂',
          recLabel: '實收偷訂',
          depLabel: '已存偷訂',
          diffLabel: '推算未入帳偷訂',
          icon: 'fa-solid fa-store text-indigo-500',
          borderAccent: 'border-t-4 border-t-indigo-500',
          bgSoft: 'bg-indigo-50/40',
          catKey: '偷訂',
          pos: 0, netRec: 0, bankDep: 0, netUnpaid: 0, fee: 0, feeRateStr: '-'
        }}
      }};

      const storesToInclude = currentStore === 'ALL' ? Object.keys(rawData.stores || {{}}) : [currentStore];
      storesToInclude.forEach(sk => {{
        const s = (rawData.stores || {{}})[sk];
        if (!s) return;
        if (s.channels) {{
          Object.entries(s.channels).forEach(([code, c]) => {{
            if (channelSummary[code]) {{
              const cPos = c.pos_sales || 0;
              const cNetRec = (c.net_after_fee !== undefined) ? c.net_after_fee : cPos;
              const cBankDep = (c.bank_deposited !== undefined) ? c.bank_deposited : (c.bank_cleared || 0);
              const cUnpaid = (c.net_unpaid_diff !== undefined) ? c.net_unpaid_diff : Math.max(0, cNetRec - cBankDep);
              const cFee = Math.max(0, cPos - cNetRec);

              channelSummary[code].pos += cPos;
              channelSummary[code].netRec += cNetRec;
              channelSummary[code].bankDep += cBankDep;
              channelSummary[code].netUnpaid += cUnpaid;
              channelSummary[code].fee += cFee;
            }}
          }});
        }}
        (s.daily_cash || []).forEach(d => {{
          if (d.audit_status === 'UNPAID') channelSummary['1130-01'].unpaidCash += (d.cash || 0);
          else if (d.audit_status === 'DECLARED_PENDING_BANK') channelSummary['1130-01'].declaredCash += (d.cash || 0);
        }});
      }});

      // 1. Render Top 4 Dedicated Channel Cards (#top-channel-cards)
      if (topCardsEl) {{
        const topCodes = ['1130-01', '1130-02', '1130-03', '1130-04'];
        const topCardsHtml = topCodes.map(code => {{
          const p = channelSummary[code];
          const pct = p.netRec > 0 ? Math.min(100, Math.round((p.bankDep / p.netRec) * 100)) : (p.bankDep > 0 ? 100 : 0);

          let statusBadge = '';
          if (p.pos === 0 && p.bankDep === 0) {{
            statusBadge = '<span class="px-2 py-0.5 rounded text-[11px] font-bold bg-slate-100 text-slate-400">本月無交易</span>';
          }} else if (p.netUnpaid <= 0) {{
            statusBadge = '<span class="px-2 py-0.5 rounded text-[11px] font-bold bg-emerald-100 text-emerald-800">✅ 全數入帳</span>';
          }} else {{
            statusBadge = `<span class="px-2 py-0.5 rounded text-[11px] font-bold bg-rose-100 text-rose-700">待入帳 $${{Math.round(p.netUnpaid).toLocaleString()}}</span>`;
          }}

          let actionFooterHtml = '';
          if (code === '1130-01') {{
            actionFooterHtml = `
              <div class="mt-3 pt-2.5 border-t border-slate-100 flex flex-wrap items-center justify-between gap-1.5 text-[11px]">
                <button onclick="jumpAndFilterAudit('UNPAID')" class="flex-1 py-1 px-2 rounded-lg bg-rose-50 hover:bg-rose-100 text-rose-700 font-bold border border-rose-200 cursor-pointer transition-all text-center">
                  🔴 未存 $${{Math.round(p.unpaidCash).toLocaleString()}}
                </button>
                <button onclick="jumpAndFilterAudit('DECLARED_PENDING_BANK')" class="flex-1 py-1 px-2 rounded-lg bg-amber-50 hover:bg-amber-100 text-amber-900 font-bold border border-amber-200 cursor-pointer transition-all text-center">
                  🟡 已報待入 $${{Math.round(p.declaredCash).toLocaleString()}}
                </button>
              </div>
            `;
          }} else {{
            actionFooterHtml = `
              <div class="mt-3 pt-2.5 border-t border-slate-100 flex items-center justify-between text-[11px]">
                <span class="text-slate-400">入帳進度 ${{pct}}%</span>
                <button onclick="jumpToBankCategory('${{p.catKey}}')" class="text-indigo-600 hover:text-indigo-800 font-bold cursor-pointer flex items-center gap-1">
                  <span>查看銀行入帳明細</span>
                  <i class="fa-solid fa-arrow-down text-[10px]"></i>
                </button>
              </div>
            `;
          }}

          const feeRowText = code === '1130-01'
            ? `<span class="text-slate-400">免平台手續費 ($0)</span>`
            : `<span class="text-amber-700 font-mono font-semibold">-$${{Math.round(p.fee).toLocaleString()}}</span>`;

          return `
            <div class="card-box p-4 flex flex-col justify-between ${{p.borderAccent}} hover:shadow-md transition-all">
              <div>
                <!-- Header -->
                <div class="flex items-center justify-between gap-2 pb-2.5 border-b border-slate-100">
                  <h3 class="font-bold text-slate-900 text-sm">${{p.cardTitle}}</h3>
                  ${{statusBadge}}
                </div>

                <!-- Primary Highlight: 實收 vs 已存 -->
                <div class="grid grid-cols-2 gap-2.5 my-3 p-2.5 rounded-xl ${{p.bgSoft}} border border-slate-200/80">
                  <div>
                    <div class="text-[11px] font-semibold text-indigo-900">${{p.recLabel}}</div>
                    <div class="text-lg font-extrabold text-indigo-700 font-mono mt-0.5">$${{Math.round(p.netRec).toLocaleString()}}</div>
                  </div>
                  <div class="border-l border-slate-200/80 pl-2.5">
                    <div class="text-[11px] font-semibold text-emerald-900">${{p.depLabel}}</div>
                    <div class="text-lg font-extrabold text-emerald-600 font-mono mt-0.5">$${{Math.round(p.bankDep).toLocaleString()}}</div>
                  </div>
                </div>

                <!-- 4-Step Deduction Breakdown -->
                <div class="space-y-1.5 text-xs">
                  <div class="flex justify-between items-center text-slate-600">
                    <span>① POS 原始營收額</span>
                    <span class="font-mono font-semibold text-slate-800">$${{Math.round(p.pos).toLocaleString()}}</span>
                  </div>
                  <div class="flex justify-between items-center text-slate-500 text-[11px]">
                    <span>➖ 平台手續費 (${{p.feeRateStr}})</span>
                    ${{feeRowText}}
                  </div>
                  <div class="flex justify-between items-center text-indigo-900 font-semibold pt-1 border-t border-slate-100">
                    <span>② 扣手續費後實收額</span>
                    <span class="font-mono font-bold text-indigo-700">$${{Math.round(p.netRec).toLocaleString()}}</span>
                  </div>
                  <div class="flex justify-between items-center text-emerald-800 font-semibold">
                    <span>③ 玉山銀行已存額</span>
                    <span class="font-mono font-bold text-emerald-600">-$${{Math.round(p.bankDep).toLocaleString()}}</span>
                  </div>
                  <div class="flex justify-between items-center pt-1.5 border-t border-slate-200 font-bold text-slate-900">
                    <span>④ 👉 ${{p.diffLabel}}</span>
                    <span class="font-mono text-sm ${{p.netUnpaid > 0 ? 'text-rose-600' : 'text-emerald-600'}}">$${{Math.round(p.netUnpaid).toLocaleString()}}</span>
                  </div>
                </div>
              </div>

              ${{actionFooterHtml}}
            </div>
          `;
        }});
        topCardsEl.innerHTML = topCardsHtml.join('');
      }}

      // 2. Render Section 4 Summary Table (#platform-cards-container)
      if (tableBodyEl) {{
        let totPos = 0, totFee = 0, totNetRec = 0, totBankDep = 0, totUnpaid = 0;
        const rowsHtml = [];

        Object.entries(channelSummary).forEach(([code, p]) => {{
          if (p.pos === 0 && p.bankDep === 0 && p.netUnpaid === 0) return;

          totPos += p.pos;
          totFee += p.fee;
          totNetRec += p.netRec;
          totBankDep += p.bankDep;
          totUnpaid += p.netUnpaid;

          let statusNoteHtml = '';
          if (p.netUnpaid <= 0) {{
            statusNoteHtml = '<span class="px-2 py-0.5 rounded bg-emerald-100 text-emerald-800 font-bold text-[11px]">✅ 全數入帳結清</span>';
          }} else if (code === '1130-01') {{
            statusNoteHtml = `<div class="flex flex-wrap items-center gap-1.5 text-[11px]">
              <span class="px-2 py-0.5 rounded bg-rose-100 text-rose-800 font-bold">🔴 門市未存 $${{Math.round(p.unpaidCash).toLocaleString()}}</span>
              <span class="px-2 py-0.5 rounded bg-amber-100 text-amber-900 font-bold">🟡 已申報待入帳 $${{Math.round(p.declaredCash).toLocaleString()}}</span>
            </div>`;
          }} else {{
            statusNoteHtml = `<span class="text-slate-600 text-[11px]">月底交易跨月撥款中（預估實收尚有 <b class="text-rose-600 font-mono">$${{Math.round(p.netUnpaid).toLocaleString()}}</b> 待入玉山存摺）</span>`;
          }}

          rowsHtml.push(`
            <tr class="hover:bg-slate-50 transition-colors">
              <td class="py-3 px-4 font-bold text-slate-900 whitespace-nowrap">
                <i class="${{p.icon}} mr-1.5"></i>${{p.name}}
              </td>
              <td class="py-3 px-4 text-right font-mono text-slate-800 whitespace-nowrap">$ ${{Math.round(p.pos).toLocaleString()}}</td>
              <td class="py-3 px-4 text-right font-mono text-amber-700 whitespace-nowrap">${{p.fee > 0 ? '-$ ' + Math.round(p.fee).toLocaleString() : '$ 0'}}</td>
              <td class="py-3 px-4 text-right font-mono font-bold text-indigo-700 bg-indigo-50/30 whitespace-nowrap">$ ${{Math.round(p.netRec).toLocaleString()}}</td>
              <td class="py-3 px-4 text-right font-mono font-bold text-emerald-600 bg-emerald-50/30 whitespace-nowrap">$ ${{Math.round(p.bankDep).toLocaleString()}}</td>
              <td class="py-3 px-4 text-right font-mono font-extrabold ${{p.netUnpaid > 0 ? 'text-rose-600 bg-rose-50/30' : 'text-slate-500'}} whitespace-nowrap">$ ${{Math.round(p.netUnpaid).toLocaleString()}}</td>
              <td class="py-3 px-4">${{statusNoteHtml}}</td>
            </tr>
          `);
        }});

        // Add Total Row
        rowsHtml.push(`
          <tr class="bg-slate-100/90 font-bold text-slate-900 border-t-2 border-slate-300">
            <td class="py-3 px-4 whitespace-nowrap">📊 全通路合計</td>
            <td class="py-3 px-4 text-right font-mono whitespace-nowrap">$ ${{Math.round(totPos).toLocaleString()}}</td>
            <td class="py-3 px-4 text-right font-mono text-amber-800 whitespace-nowrap">-$ ${{Math.round(totFee).toLocaleString()}}</td>
            <td class="py-3 px-4 text-right font-mono text-indigo-700 whitespace-nowrap">$ ${{Math.round(totNetRec).toLocaleString()}}</td>
            <td class="py-3 px-4 text-right font-mono text-emerald-700 whitespace-nowrap">$ ${{Math.round(totBankDep).toLocaleString()}}</td>
            <td class="py-3 px-4 text-right font-mono text-rose-600 whitespace-nowrap">$ ${{Math.round(totUnpaid).toLocaleString()}}</td>
            <td class="py-3 px-4 text-slate-600 text-[11px]">實收總額 $${{Math.round(totNetRec).toLocaleString()}} - 已存總額 $${{Math.round(totBankDep).toLocaleString()}} = 未入帳總差額 $${{Math.round(totUnpaid).toLocaleString()}}</td>
          </tr>
        `);

        tableBodyEl.innerHTML = rowsHtml.join('');
      }}
    }}

    function renderInvoiceTracking(rawData) {{
      const tbody = document.getElementById('invoice-tracking-tbody');
      if (!tbody) return;

      const storesToInclude = currentStore === 'ALL' ? Object.keys(rawData.stores || {{}}) : [currentStore];
      let rows = [];
      let totalVat = 0;
      let totalFee = 0;

      storesToInclude.forEach(sk => {{
        const s = (rawData.stores || {{}})[sk];
        if (!s) return;
        (s.invoice_tracking || []).forEach(it => {{
          rows.push(it);
          totalVat += (it.est_vat_5pct || 0);
          totalFee += (it.est_total_fee || 0);
        }});
      }});

      const summaryBadge = document.getElementById('invoice-summary-badge');
      if (summaryBadge) {{
        summaryBadge.textContent = `預估全月平台費 $${{Math.round(totalFee).toLocaleString()}} ｜ 可扣抵 5% 進項稅額合計：$${{Math.round(totalVat).toLocaleString()}}`;
      }}

      if (rows.length === 0) {{
        tbody.innerHTML = `<tr><td colspan="9" class="py-6 text-center text-slate-400">本月無外送或電子支付平台手續費紀錄</td></tr>`;
        return;
      }}

      tbody.innerHTML = rows.map(it => {{
        const sClean = (it.store || '').replace(' (月下餐飲)', '');
        const isYellow = (it.status || '').includes('🟡');
        const isGreen = (it.status || '').includes('🟢') || (it.status || '').includes('✅');
        const stBadge = isGreen
          ? `<span class="inline-block px-2.5 py-1 rounded-full bg-emerald-100 text-emerald-800 font-bold text-[11px]">${{it.status}}</span>`
          : isYellow
          ? `<span class="inline-block px-2.5 py-1 rounded-full bg-amber-100 text-amber-900 font-bold text-[11px]">${{it.status}}</span>`
          : `<span class="inline-block px-2.5 py-1 rounded-full bg-rose-100 text-rose-800 font-bold text-[11px]">${{it.status}}</span>`;

        const regInvs = it.registered_invoices || [];
        const regHtml = regInvs.length > 0
          ? regInvs.map(r => `
              <div class="inline-flex items-center gap-1.5 px-2 py-0.5 rounded-lg bg-emerald-50 text-emerald-800 border border-emerald-200 text-[11px] font-mono mr-1 mb-1">
                <i class="fa-solid fa-receipt text-emerald-600"></i>
                <span class="font-bold">${{r.invoice_no}}</span>
                <span>($${{Math.round(r.amount_total || 0).toLocaleString()}})</span>
                ${{r.file_path ? `<a href="${{r.file_path}}" target="_blank" title="下載/檢視發票檔案" class="text-blue-600 hover:text-blue-800 ml-0.5"><i class="fa-solid fa-paperclip"></i></a>` : ''}}
                <button onclick="deleteInvoiceRecord('${{r.id}}')" title="刪除此發票紀錄" class="text-rose-400 hover:text-rose-700 ml-1 cursor-pointer font-bold">✕</button>
              </div>
            `).join('')
          : `<span class="text-slate-400 italic text-[11px]">尚無登記</span>`;

        const btnLabel = regInvs.length > 0 ? '➕ 追加發票' : '📤 登記發票';
        const estFillAmt = Math.round(it.unbilled_fee > 0 ? it.unbilled_fee : it.est_total_fee);

        return `
          <tr class="hover:bg-slate-50 transition-colors">
            <td class="py-3 px-4 font-bold text-slate-900 whitespace-nowrap">${{sClean}}</td>
            <td class="py-3 px-4 whitespace-nowrap">
              <span class="font-bold text-slate-800">${{it.platform}}</span>
              <span class="ml-1.5 px-1.5 py-0.5 rounded bg-slate-100 text-slate-600 font-mono text-[10px]">統編 ${{it.tax_id}}</span>
            </td>
            <td class="py-3 px-4 text-center font-mono text-slate-600 whitespace-nowrap">${{it.fee_rate}}</td>
            <td class="py-3 px-4 text-right font-mono text-slate-700 whitespace-nowrap">$ ${{Math.round(it.pos_sales || 0).toLocaleString()}}</td>
            <td class="py-3 px-4 text-right font-mono font-bold text-slate-900 whitespace-nowrap">$ ${{Math.round(it.est_total_fee || 0).toLocaleString()}}</td>
            <td class="py-3 px-4 text-right font-mono font-bold text-blue-600 whitespace-nowrap">$ ${{Math.round(it.est_vat_5pct || 0).toLocaleString()}}</td>
            <td class="py-3 px-4 text-center whitespace-nowrap">${{stBadge}}</td>
            <td class="py-3 px-4">${{regHtml}}</td>
            <td class="py-3 px-4">
              <div class="flex items-center gap-2">
                <button onclick="openInvoiceUploadModal('${{sClean}}', '${{it.platform}}', ${{estFillAmt}})" 
                        class="px-2.5 py-1 rounded-lg bg-indigo-50 hover:bg-indigo-100 text-indigo-700 font-bold border border-indigo-200 text-[11px] whitespace-nowrap cursor-pointer transition-all inline-flex items-center gap-1 shadow-sm">
                  <span>${{btnLabel}}</span>
                </button>
                <span class="text-[11px] text-slate-500">${{it.action}}</span>
              </div>
            </td>
          </tr>
        `;
      }}).join('');
    }}

    function showCalendarDayDetail(dateStr) {{
      const detailBox = document.getElementById('calendar-day-detail');
      if (!detailBox) return;
      const records = currentDailyCashMap[dateStr] || [];
      if (records.length === 0) {{
        detailBox.classList.remove('hidden');
        detailBox.innerHTML = `<div class="flex justify-between items-center"><span>📅 <b>${{dateStr}}</b>：當日無門市現金營收</span><button onclick="document.getElementById('calendar-day-detail').classList.add('hidden')" class="text-slate-300 hover:text-white cursor-pointer">✕</button></div>`;
        return;
      }}
      const itemsHtml = records.map(r => {{
        const t1 = r.has_declared ? `✅ 已申報 (${{(r.declared_note || '').replace('申報於 ', '')}})` : '❌ 總表U欄空白未填';
        const t2 = r.bank_cleared ? `🏦 玉山 ${{r.bank_date || r.deposit_date || ''}} 入帳` : '⏳ 銀行存摺尚未入帳';
        return `<div class="mt-2 pt-2 border-t border-slate-700 grid grid-cols-1 sm:grid-cols-4 gap-2 items-center">
          <span class="font-bold text-white">${{r.store.replace(' (月下餐飲)', '')}} · <span class="text-amber-300 font-mono">$${{Math.round(r.cash).toLocaleString()}}</span></span>
          <span class="text-slate-200">第一層：${{t1}}</span>
          <span class="text-slate-200">第二層：${{t2}}</span>
          <span class="sm:text-right font-bold text-white">${{r.audit_label || ''}}</span>
        </div>`;
      }}).join('');
      detailBox.classList.remove('hidden');
      detailBox.innerHTML = `
        <div class="flex justify-between items-center font-bold text-indigo-300">
          <span>📅 ${{dateStr}} 雙層審核詳情</span>
          <button onclick="document.getElementById('calendar-day-detail').classList.add('hidden')" class="text-slate-300 hover:text-white px-2 py-0.5 rounded bg-slate-700 cursor-pointer">✕ 關閉</button>
        </div>
        ${{itemsHtml}}
      `;
    }}

    function renderCalendar(rawData) {{
      const grid = document.getElementById('calendar-grid');
      if (!grid) return;
      
      const [yr, mo] = currentMonth.split('-').map(Number);
      document.getElementById('calendar-title-text').textContent = 
        `營業與存款核對月曆 (${{yr}} 年 ${{mo}} 月)`;

      const daysInMonth = new Date(yr, mo, 0).getDate();
      const firstDayOfWeek = new Date(yr, mo - 1, 1).getDay();

      let cellsHtml = [];
      for (let i = 0; i < firstDayOfWeek; i++) {{
        cellsHtml.push('<div class="calendar-cell bg-slate-50/40 rounded-xl border border-dashed border-slate-200"></div>');
      }}

      let dailyCashMap = {{}};
      const storesToInclude = currentStore === 'ALL' ? Object.keys(rawData.stores || {{}}) : [currentStore];
      
      storesToInclude.forEach(sk => {{
        const s = (rawData.stores || {{}})[sk];
        if (s && s.daily_cash) {{
          s.daily_cash.forEach(d => {{
            if (!dailyCashMap[d.date]) {{
              dailyCashMap[d.date] = [];
            }}
            dailyCashMap[d.date].push({{ store: s.name, ...d }});
          }});
        }}
      }});
      currentDailyCashMap = dailyCashMap;

      for (let day = 1; day <= daysInMonth; day++) {{
        const dayStr = day < 10 ? '0' + day : '' + day;
        const moStr = mo < 10 ? '0' + mo : '' + mo;
        const dateStr = `${{yr}}-${{moStr}}-${{dayStr}}`;
        const dayRecords = dailyCashMap[dateStr] || [];

        let innerBadges = '';
        if (dayRecords.length > 0) {{
          dayRecords.forEach(r => {{
            if (r.cash > 0) {{
              const sShort = r.store.slice(0, 2);
              const amtStr = '$' + Math.round(r.cash).toLocaleString();
              if (r.audit_status === 'MATCHED') {{
                innerBadges += `<div class="px-1.5 py-0.5 rounded text-[11px] font-semibold mt-1 flex justify-between items-center bg-emerald-50 text-emerald-800 border border-emerald-200">
                  <span>🟢 ${{sShort}}</span>
                  <span class="font-mono">${{amtStr}}</span>
                </div>`;
              }} else if (r.audit_status === 'DECLARED_PENDING_BANK') {{
                innerBadges += `<div class="px-1.5 py-0.5 rounded text-[11px] font-bold mt-1 flex justify-between items-center bg-amber-50 text-amber-900 border border-amber-300">
                  <span>🟡 ${{sShort}}</span>
                  <span class="font-mono">${{amtStr}}</span>
                </div>`;
              }} else if (r.audit_status === 'BANK_ONLY') {{
                innerBadges += `<div class="px-1.5 py-0.5 rounded text-[11px] font-bold mt-1 flex justify-between items-center bg-blue-50 text-blue-800 border border-blue-200">
                  <span>🔵 ${{sShort}}</span>
                  <span class="font-mono">${{amtStr}}</span>
                </div>`;
              }} else {{
                innerBadges += `<div class="px-1.5 py-0.5 rounded text-[11px] font-bold mt-1 flex justify-between items-center bg-rose-50 text-rose-800 border border-rose-200">
                  <span>🔴 ${{sShort}}</span>
                  <span class="font-mono">${{amtStr}}</span>
                </div>`;
              }}
            }}
          }});
        }}

        cellsHtml.push(`
          <div id="cal-date-${{dateStr}}" onclick="showCalendarDayDetail('${{dateStr}}')" class="calendar-cell bg-white p-2 rounded-xl border border-slate-200 flex flex-col justify-between text-xs cursor-pointer">
            <div>
              <div class="flex justify-between items-center mb-1">
                <span class="font-bold text-slate-700">${{day}}</span>
                <span class="text-[10px] text-slate-400">${{mo}}/${{dayStr}}</span>
              </div>
              <div>
                ${{innerBadges || '<div class="text-[10px] text-slate-300 mt-3 text-center">-</div>'}}
              </div>
            </div>
          </div>
        `);
      }}
      grid.innerHTML = cellsHtml.join('');
    }}

    function filterTransactions(rawDataArg) {{
      const rawData = rawDataArg || allMonthsData[currentMonth];
      if (!rawData) return;

      const tbody = document.getElementById('transaction-tbody');
      if (!tbody) return;
      const query = ((document.getElementById('search-input') || {{}}).value || '').toLowerCase();

      const filtered = (rawData.cleared_items || []).filter(item => {{
        if (currentStore !== 'ALL' && !item.store.includes(currentStore)) return false;
        if (currentCategoryFilter === 'CROSS') {{
          const isCross = item.is_cross_month || (item.revenue_alloc && (item.revenue_alloc.includes('跨月') || item.revenue_alloc.includes('上月') || item.revenue_alloc.includes('拆分')));
          if (!isCross) return false;
        }} else if (currentCategoryFilter !== 'ALL') {{
          if (!item.category.includes(currentCategoryFilter)) return false;
        }}
        if (query) {{
          const str = (item.store + ' ' + item.bank_date + ' ' + item.category + ' ' + item.period + ' ' + (item.revenue_alloc || '') + ' ' + item.remark + ' ' + item.bank_amt).toLowerCase();
          return str.includes(query);
        }}
        return true;
      }});

      if (filtered.length === 0) {{
        tbody.innerHTML = `<tr><td colspan="7" class="py-8 text-center text-slate-400">查無相符之單筆沖銷明細</td></tr>`;
        return;
      }}

      const rowsHtml = filtered.map(item => {{
        const isCash = item.category.includes('現金');
        const isLine = item.category.includes('LINE');
        const isUber = item.category.includes('Uber');
        const isPanda = item.category.includes('Foodpanda');
        const isTouding = item.category.includes('偷訂');

        const catBadge = isCash ? '<span class="px-2 py-0.5 rounded bg-amber-100 text-amber-900 font-bold text-[11px]">💵 現金存入</span>' :
          isLine ? '<span class="px-2 py-0.5 rounded bg-emerald-100 text-emerald-800 font-bold text-[11px]">💚 LINE Pay</span>' :
          isUber ? '<span class="px-2 py-0.5 rounded bg-slate-800 text-white font-bold text-[11px]">🚗 Uber Eats</span>' :
          isPanda ? '<span class="px-2 py-0.5 rounded bg-pink-100 text-pink-800 font-bold text-[11px]">🐼 Foodpanda</span>' :
          isTouding ? '<span class="px-2 py-0.5 rounded bg-purple-100 text-purple-800 font-bold text-[11px]">🛒 偷訂</span>' :
          `<span class="px-2 py-0.5 rounded bg-slate-100 text-slate-700 font-bold text-[11px]">${{item.category}}</span>`;

        const cleanPeriod = (item.period || '').replace(/2026-/g, '');
        const alloc = item.revenue_alloc || '全數歸本月';
        const isCross = alloc.includes('跨月') || alloc.includes('上月') || alloc.includes('拆分');
        const allocHtml = isCross
          ? `<span class="text-amber-700 font-semibold">${{alloc}}</span>`
          : `<span class="text-slate-500">${{alloc}}</span>`;

        return `
          <tr class="hover:bg-slate-50 transition-colors">
            <td class="py-2.5 px-4 font-mono text-slate-400 text-center">${{item.idx}}</td>
            <td class="py-2.5 px-4 font-bold text-slate-800 whitespace-nowrap">${{item.store.replace(' (月下餐飲)', '')}}</td>
            <td class="py-2.5 px-4 font-mono text-slate-600 whitespace-nowrap">${{item.bank_date.slice(5)}}</td>
            <td class="py-2.5 px-4 whitespace-nowrap">${{catBadge}}</td>
            <td class="py-2.5 px-4 text-right font-mono font-bold text-slate-900 whitespace-nowrap">$ ${{Math.round(item.bank_amt).toLocaleString()}}</td>
            <td class="py-2.5 px-4 font-mono text-slate-700 font-medium">${{cleanPeriod}}</td>
            <td class="py-2.5 px-4">${{allocHtml}}</td>
          </tr>
        `;
      }});
      tbody.innerHTML = rowsHtml.join('');
    }}

    function scrollToDate(dateStr) {{
      switchMainTab('cash');
      toggleCashView('calendar');
      showCalendarDayDetail(dateStr);
      const el = document.getElementById('cal-date-' + dateStr);
      if (el) {{
        el.scrollIntoView({{ behavior: 'smooth', block: 'center' }});
        el.classList.add('ring-2', 'ring-indigo-600');
        setTimeout(() => el.classList.remove('ring-2', 'ring-indigo-600'), 2500);
      }}
    }}

    function getApiHost() {{
      if (typeof window !== 'undefined' && window.location && (window.location.hostname === 'localhost' || window.location.hostname === '127.0.0.1')) {{
        return '';
      }}
      return 'http://localhost:8080';
    }}

    function openInvoiceUploadModal(store = '', platform = '', estAmount = 0) {{
      const modal = document.getElementById('invoice-upload-modal');
      if (!modal) return;

      const storeSelect = document.getElementById('modal-inv-store');
      const platformSelect = document.getElementById('modal-inv-platform');
      const noInput = document.getElementById('modal-inv-no');
      const dateInput = document.getElementById('modal-inv-date');
      const amountInput = document.getElementById('modal-inv-amount');
      const noteInput = document.getElementById('modal-inv-note');
      const errorEl = document.getElementById('modal-error-msg');

      if (errorEl) errorEl.classList.add('hidden');
      if (noInput) noInput.value = '';
      if (noteInput) noteInput.value = '';
      clearSelectedFile();

      if (storeSelect) {{
        if (store && store.includes('隨意')) storeSelect.value = '隨意小吃';
        else storeSelect.value = '月宮';
      }}

      if (platformSelect) {{
        if (platform) {{
          if (platform.includes('LINE') || platform.includes('連加')) platformSelect.value = 'LINE Pay (連加網路)';
          else if (platform.includes('Uber') || platform.includes('優步')) platformSelect.value = 'Uber Eats (優步科技)';
          else if (platform.includes('Foodpanda') || platform.includes('富胖達')) platformSelect.value = 'Foodpanda (富胖達)';
        }}
      }}

      if (dateInput) {{
        const today = new Date().toISOString().split('T')[0];
        dateInput.value = today;
      }}

      if (amountInput) {{
        if (estAmount && estAmount > 0) {{
          amountInput.value = Math.round(estAmount);
        }} else {{
          amountInput.value = '';
        }}
        calculateVatSplit();
      }}

      modal.classList.remove('hidden');
    }}

    function closeInvoiceUploadModal() {{
      const modal = document.getElementById('invoice-upload-modal');
      if (modal) modal.classList.add('hidden');
    }}

    function calculateVatSplit() {{
      const amountInput = document.getElementById('modal-inv-amount');
      const untaxedEl = document.getElementById('modal-inv-untaxed');
      const taxEl = document.getElementById('modal-inv-tax');
      const val = parseFloat(amountInput ? amountInput.value : 0) || 0;
      const untaxed = Math.round(val / 1.05);
      const vat = Math.round(val - untaxed);
      if (untaxedEl) untaxedEl.textContent = '$' + untaxed.toLocaleString();
      if (taxEl) taxEl.textContent = '$' + vat.toLocaleString();
    }}

    function handleFileSelected(input) {{
      const textEl = document.getElementById('drop-zone-text');
      const selEl = document.getElementById('drop-zone-selected');
      const nameEl = document.getElementById('selected-file-name');
      if (input.files && input.files[0]) {{
        if (nameEl) nameEl.textContent = input.files[0].name;
        if (textEl) textEl.classList.add('hidden');
        if (selEl) {{
          selEl.classList.remove('hidden');
          selEl.classList.add('flex');
        }}
      }}
    }}

    function clearSelectedFile() {{
      const input = document.getElementById('modal-inv-file');
      const textEl = document.getElementById('drop-zone-text');
      const selEl = document.getElementById('drop-zone-selected');
      if (input) input.value = '';
      if (textEl) textEl.classList.remove('hidden');
      if (selEl) {{
        selEl.classList.add('hidden');
        selEl.classList.remove('flex');
      }}
    }}

    function setupDropZone() {{
      const dropZone = document.getElementById('drop-zone');
      if (!dropZone) return;
      ['dragenter', 'dragover'].forEach(eventName => {{
        dropZone.addEventListener(eventName, (e) => {{
          e.preventDefault();
          e.stopPropagation();
          dropZone.classList.add('border-indigo-500', 'bg-indigo-50/50');
        }}, false);
      }});
      ['dragleave', 'drop'].forEach(eventName => {{
        dropZone.addEventListener(eventName, (e) => {{
          e.preventDefault();
          e.stopPropagation();
          dropZone.classList.remove('border-indigo-500', 'bg-indigo-50/50');
        }}, false);
      }});
      dropZone.addEventListener('drop', (e) => {{
        const dt = e.dataTransfer;
        const files = dt.files;
        const fileInput = document.getElementById('modal-inv-file');
        if (fileInput && files.length > 0) {{
          fileInput.files = files;
          handleFileSelected(fileInput);
        }}
      }}, false);
    }}

    async function handleInvoiceFormSubmit(event) {{
      event.preventDefault();
      const errorEl = document.getElementById('modal-error-msg');
      if (errorEl) errorEl.classList.add('hidden');

      const submitBtn = document.getElementById('modal-submit-btn');
      const origBtnHtml = submitBtn ? submitBtn.innerHTML : '';
      if (submitBtn) {{
        submitBtn.disabled = true;
        submitBtn.innerHTML = '<i class="fa-solid fa-spinner fa-spin"></i> <span>核銷計算中...</span>';
      }}

      const store = document.getElementById('modal-inv-store').value;
      const platform = document.getElementById('modal-inv-platform').value;
      const invNo = document.getElementById('modal-inv-no').value.trim();
      const invDate = document.getElementById('modal-inv-date').value;
      const amountTotal = parseFloat(document.getElementById('modal-inv-amount').value) || 0;
      const note = document.getElementById('modal-inv-note').value.trim();
      const fileInput = document.getElementById('modal-inv-file');

      let fileData = null;
      let fileName = null;

      if (fileInput && fileInput.files && fileInput.files.length > 0) {{
        const file = fileInput.files[0];
        fileName = file.name;
        try {{
          fileData = await new Promise((resolve, reject) => {{
            const reader = new FileReader();
            reader.onload = () => resolve(reader.result);
            reader.onerror = reject;
            reader.readAsDataURL(file);
          }});
        }} catch (fe) {{
          console.warn("File read error:", fe);
        }}
      }}

      const payload = {{
        month: currentMonth,
        store: store,
        platform: platform,
        invoice_no: invNo,
        invoice_date: invDate,
        amount_total: amountTotal,
        note: note,
        file_name: fileName,
        file_data: fileData
      }};

      try {{
        const apiHost = getApiHost();
        const res = await fetch(apiHost + '/api/invoices/upload', {{
          method: 'POST',
          headers: {{ 'Content-Type': 'application/json' }},
          body: JSON.stringify(payload)
        }});
        const result = await res.json();
        if (res.ok && result && result.success) {{
          if (result.data) {{
            allMonthsData = result.data;
            renderDashboard();
          }}
          closeInvoiceUploadModal();
          showToast(`✅ 發票 ${{invNo}} 已登記並成功抵減手續費！`);
        }} else {{
          throw new Error(result.error || '登記失敗，請檢查資料');
        }}
      }} catch (err) {{
        if (errorEl) {{
          errorEl.textContent = '⚠️ 發票儲存失敗：' + (err.message || '無法連線至本機伺服器，請確認伺服器已啟動');
          errorEl.classList.remove('hidden');
        }}
      }} finally {{
        if (submitBtn) {{
          submitBtn.disabled = false;
          submitBtn.innerHTML = origBtnHtml;
        }}
      }}
    }}

    async function deleteInvoiceRecord(id) {{
      if (!confirm('確定要刪除此筆發票登記紀錄嗎？\\n系統將自動復原對帳差額與營業稅試算。')) {{
        return;
      }}
      try {{
        const apiHost = getApiHost();
        const res = await fetch(apiHost + '/api/invoices/delete', {{
          method: 'POST',
          headers: {{ 'Content-Type': 'application/json' }},
          body: JSON.stringify({{ month: currentMonth, id: id }})
        }});
        const result = await res.json();
        if (res.ok && result && result.success) {{
          if (result.data) {{
            allMonthsData = result.data;
            renderDashboard();
          }}
          showToast('🗑️ 發票紀錄已刪除，對帳差額已重新試算！');
        }} else {{
          alert('刪除失敗：' + (result.error || '請稍候重試'));
        }}
      }} catch (err) {{
        alert('無法連線本機服務進行刪除：' + err.message);
      }}
    }}

    setupDropZone();
    renderDashboard();
    refreshLiveData(false);
  </script>
</body>
</html>
"""

    out_file = os.path.join(OUTPUT_DIR, f"對帳看板_{m_clean}.html")
    with open(out_file, "w", encoding="utf-8") as f:
        f.write(html_content)

    latest_file = os.path.join(OUTPUT_DIR, "最新對帳看板.html")
    with open(latest_file, "w", encoding="utf-8") as f:
        f.write(html_content)

    root_index = os.path.join(BASE_DIR, "index.html")
    with open(root_index, "w", encoding="utf-8") as f:
        f.write(html_content)

    docs_dir = os.path.join(BASE_DIR, "docs")
    os.makedirs(docs_dir, exist_ok=True)
    docs_index = os.path.join(docs_dir, "index.html")
    with open(docs_index, "w", encoding="utf-8") as f:
        f.write(html_content)

    # 同步複製發票紀錄與附件至 docs/ 供靜態看板讀取
    inv_rec_file = os.path.join(BASE_DIR, "invoices_record.json")
    if os.path.exists(inv_rec_file):
        shutil.copy2(inv_rec_file, os.path.join(docs_dir, "invoices_record.json"))

    docs_inv_dir = os.path.join(docs_dir, "uploaded_invoices")
    os.makedirs(docs_inv_dir, exist_ok=True)
    up_inv_dir = os.path.join(BASE_DIR, "uploaded_invoices")
    if os.path.exists(up_inv_dir):
        for fname in os.listdir(up_inv_dir):
            src_f = os.path.join(up_inv_dir, fname)
            if os.path.isfile(src_f):
                shutil.copy2(src_f, os.path.join(docs_inv_dir, fname))

    # 同步複製 Excel 報表與傳票至 docs/output 供 GitHub Pages 跨裝置下載
    docs_out_dir = os.path.join(docs_dir, "output")
    os.makedirs(docs_out_dir, exist_ok=True)
    for fname in os.listdir(OUTPUT_DIR):
        if fname.endswith(".xlsx"):
            shutil.copy2(os.path.join(OUTPUT_DIR, fname), os.path.join(docs_out_dir, fname))

    print(f"🌐 可視化查帳互動看板已產出：{out_file}")
    print(f"👉 點擊開啟最新看板：{latest_file}")
    print(f"🚀 GitHub Pages 網頁入口已更新：{root_index}")

    return out_file

if __name__ == "__main__":
    month = sys.argv[1] if len(sys.argv) > 1 else "2026-10"
    generate_html_dashboard(target_month=month)
